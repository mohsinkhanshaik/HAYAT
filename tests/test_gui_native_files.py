"""Direct native edits run in workers and preserve the original on conflicts."""
import queue
import time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from types import SimpleNamespace

import pytest

from app.gui_files import FileEditor, NativeFileEditor, open_editor
from app.gui_theme import Palette
from app.safety.permissions import PermissionManager
from app.tools.filesystem import register_filesystem_tools
from app.tools.registry import ToolRegistry


@pytest.fixture
def gui(shared_tk, tmp_path):
    class GUI(Palette):
        def _post(self, function, *args):
            self.events.put((function, args))
    window = GUI()
    window.root = shared_tk
    window.events = queue.SimpleQueue()
    window._preview_workers = ThreadPoolExecutor(max_workers=2)
    registry = ToolRegistry(PermissionManager({'default': 'ALLOW'}, tmp_path))
    filesystem = register_filesystem_tools(registry)
    window.runtime = SimpleNamespace(registry=registry)
    window.filesystem = filesystem
    yield window
    window._preview_workers.shutdown(wait=True)


def wait(gui, predicate):
    deadline = time.monotonic() + 5
    while time.monotonic() < deadline:
        gui.root.update()
        while not gui.events.empty():
            callback, arguments = gui.events.get_nowait()
            callback(*arguments)
        if predicate():
            return
        time.sleep(.01)
    assert predicate()


@pytest.mark.parametrize('suffix', ['.md', '.docx'])
@pytest.mark.parametrize('failure', ['worker', 'submission'])
def test_an_editor_failure_keeps_edits_and_allows_a_retry(gui, tmp_path, monkeypatch, suffix, failure):
    path = tmp_path / ('recoverable' + suffix)
    if suffix == '.docx':
        gui.filesystem.write_docx(str(path), 'Before')
    else:
        path.write_text('Before', encoding='utf-8')
    original = path.read_bytes()
    editor = open_editor(gui, path)
    try:
        wait(gui, lambda: editor._digest and not editor._busy)
        if isinstance(editor, FileEditor):
            editor.text.delete('1.0', 'end')
            editor.text.insert('1.0', 'After')
        else:
            editor.old_text.insert('1.0', 'Before')
            editor.new_text.insert('1.0', 'After')

        def fail(*args, **kwargs):
            raise RuntimeError('Editor operation interrupted')

        with monkeypatch.context() as patch:
            target = gui.runtime.registry if failure == 'worker' else gui._preview_workers
            patch.setattr(target, 'execute' if failure == 'worker' else 'submit', fail)
            editor.save()
            wait(gui, lambda: not editor._busy)
        assert path.read_bytes() == original
        assert 'Editor operation interrupted' in editor.status.get()
        assert str(editor.save_button.cget('state')) == 'normal'
        if isinstance(editor, FileEditor):
            assert str(editor.text.cget('state')) == 'normal'
            assert editor.text.get('1.0', 'end-1c') == 'After'
        else:
            assert str(editor.new_text.cget('state')) == 'normal'
            assert editor.new_text.get('1.0', 'end-1c') == 'After'
        editor.save()
        wait(gui, lambda: not editor._busy)
        assert path.read_bytes() != original
        assert list(tmp_path.glob(path.name + '.hayat-*.bak'))
    finally:
        editor.destroy()


@pytest.mark.parametrize('suffix', ['.docx', '.pptx'])
def test_document_editor_applies_an_edit_and_rejects_external_changes(gui, tmp_path, suffix):
    from docx import Document
    from pptx import Presentation
    path = tmp_path / ('native' + suffix)
    if suffix == '.docx':
        gui.filesystem.write_docx(str(path), 'Before and untouched')
    else:
        from app.tools.file_editing import FileEditing
        FileEditing(gui.filesystem).write_presentation(str(path), [{'title': 'Before', 'body': 'untouched'}])
    editor = open_editor(gui, path)
    try:
        assert isinstance(editor, NativeFileEditor)
        wait(gui, lambda: editor._digest and not editor._busy)
        editor.old_text.insert('1.0', 'Before')
        editor.new_text.insert('1.0', 'After')
        editor.save()
        wait(gui, lambda: editor._last_backup and not editor._busy)
        content = Document(path).paragraphs[0].text if suffix == '.docx' else Presentation(path).slides[0].shapes.title.text
        assert 'After' in content
        assert Path(editor._last_backup).is_file()
        if suffix == '.docx':
            gui.filesystem.write_docx(str(path), 'External change')
        else:
            FileEditing(gui.filesystem).write_presentation(str(path), [{'title': 'External change', 'body': 'untouched'}])
        before = path.read_bytes()
        editor.save()
        wait(gui, lambda: not editor._busy)
        assert 'changed' in editor.status.get()
        assert path.read_bytes() == before
    finally:
        editor.destroy()


def test_spreadsheet_editor_changes_a_typed_cell_and_retains_other_formulas(gui, tmp_path):
    import openpyxl
    from app.tools.file_editing import FileEditing
    path = tmp_path / 'native.xlsx'
    FileEditing(gui.filesystem).write_spreadsheet(str(path), [['Task', '=1+1'], ['Original', 2]])
    editor = open_editor(gui, path)
    try:
        wait(gui, lambda: editor._digest and not editor._busy)
        assert 'Sheet1' in editor.sheet_picker.cget('values')
        editor.values['cell'].set('A2')
        editor.values['value'].set('=literal')
        editor.save()
        wait(gui, lambda: editor._last_backup and not editor._busy)
        book = openpyxl.load_workbook(path)
        assert book.active['A2'].value == '=literal' and book.active['A2'].data_type == 's'
        assert book.active['B1'].value == '=1+1'
        book.close()
        editor.values['type'].set('Number')
        editor.values['value'].set('NaN')
        editor.save()
        assert 'finite' in editor.status.get() and not editor._busy
    finally:
        editor.destroy()
