"""Persistent draft editing, native starter files and delivery verification."""
import zipfile
from pathlib import Path

import pytest

from app.safety.permissions import PermissionManager
from app.tools.business import BusinessLedger
from app.tools.filesystem import register_filesystem_tools
from app.tools.product_workshop import register_product_workspace_tools
from app.tools.products import ProductStore, register_product_tools
from app.tools.registry import ToolRegistry


@pytest.fixture
def workshop(tmp_path):
    permission = PermissionManager({'default': 'ALLOW'}, tmp_path)
    registry = ToolRegistry(permission)
    filesystem = register_filesystem_tools(registry)
    store = ProductStore(tmp_path / 'products.db', tmp_path / 'artifacts/products', permission)
    business = BusinessLedger(tmp_path / 'business.db', tmp_path / 'artifacts/business', permission)
    business.set_profile(business_name='Hamuzair', owner_name='Mohsin Khan Shaik')
    register_product_tools(registry, store)
    register_product_workspace_tools(registry, store, filesystem, business)
    product = store.create('Useful weekly kit', 'Independent creators', 'Plan useful work with actual evidence.', '9.99', 'Personal use only.')
    return store, registry, product


@pytest.mark.parametrize('kind', ['weekly_planner', 'content_calendar', 'research_report'])
def test_starters_make_real_branded_native_files_and_verified_delivery(workshop, kind):
    import openpyxl
    from docx import Document
    from pptx import Presentation
    from pypdf import PdfReader

    store, registry, product = workshop
    result = registry.execute('product_build_starter', {'product_id': product['id'], 'kind': kind})
    assert result.success, result.error
    paths = {Path(asset['path']).suffix: Path(asset['path']) for asset in result.output['assets']}
    assert set(paths) == {'.md', '.pdf', '.docx', '.xlsx', '.pptx'}
    assert 'Hamuzair' in paths['.md'].read_text(encoding='utf-8')
    assert any('Mohsin Khan Shaik' in paragraph.text for paragraph in Document(paths['.docx']).paragraphs)
    assert len(PdfReader(paths['.pdf']).pages) > 0
    assert len(Presentation(paths['.pptx']).slides) == 4
    workbook = openpyxl.load_workbook(paths['.xlsx'])
    assert len(workbook.sheetnames) == 3
    assert workbook['Overview']['B6'].data_type == 'f'
    assert len(workbook.worksheets[1].data_validations.dataValidation) == 1
    workbook.close()
    package = store.package(product['id'], [asset['path'] for asset in result.output['assets']])
    checked = store.verify_bundle(product['id'])
    assert checked['verified'] and not checked['needs_rebuild']
    reopened = ProductStore(store.path, store.output, store.permissions)
    assert reopened.details(product['id'])['bundle']['path'] == package['path']
    assert len(reopened.details(product['id'])['assets']) == 5


def test_edit_conflicts_and_sale_currency_history_are_preserved(workshop):
    store, registry, product = workshop
    old = store.details(product['id'])
    arguments = {key: product[key] for key in ('title', 'audience', 'description', 'license_text', 'currency')}
    arguments.update(product_id=product['id'], price='12.00', expected_sha256=old['metadata_sha256'])
    saved = registry.execute('product_update', arguments)
    assert saved.success and saved.output['product']['price_cents'] == 1200
    conflict = registry.execute('product_update', arguments)
    assert not conflict.success and 'changed since' in conflict.error
    store.record_sale(product['id'], 'actual-reference', '9.99')
    arguments.update(currency='EUR', expected_sha256=saved.output['metadata_sha256'])
    changed = registry.execute('product_update', arguments)
    assert not changed.success and 'Currency cannot change' in changed.error
    assert store.report()['by_currency'][0]['currency'] == 'USD'


def test_delivery_detects_stale_sources_listing_and_tampering(workshop, tmp_path):
    store, registry, product = workshop
    source = tmp_path / 'template.md'
    source.write_text('Version one', encoding='utf-8')
    package = store.package(product['id'], [str(source)])
    source.write_text('Version two', encoding='utf-8')
    checked = store.verify_bundle(product['id'])
    assert checked['verified'] and checked['needs_rebuild']
    assert checked['sources_changed'] == ['template.md']
    old = store.details(product['id'])
    store.update(product['id'], 'Revised title', product['audience'], product['description'], '9.99', product['license_text'], old['metadata_sha256'])
    assert store.verify_bundle(product['id'])['listing_changed']
    with zipfile.ZipFile(package['path'], 'a') as archive:
        archive.writestr('unreviewed.txt', 'Extra file')
    result = registry.execute('product_verify_bundle', {'product_id': product['id']})
    assert result.success and not result.output['verified']
    assert result.output['errors']
    assert store.catalog()['status'].startswith('local drafts')


def test_asset_validation_is_atomic_and_empty_selection_is_supported(workshop, tmp_path):
    store, _, product = workshop
    source = tmp_path / 'deliverable.md'
    source.write_text('A real file')
    store.set_files(product['id'], [str(source)])
    with pytest.raises(ValueError):
        store.set_files(product['id'], [str(tmp_path / 'missing.md')])
    assert store.details(product['id'])['assets'][0]['path'] == str(source)
    assert store.set_files(product['id'], [])['assets'] == []


def test_new_product_tools_enforce_permissions(workshop, tmp_path):
    store, _, product = workshop
    denied = ToolRegistry(PermissionManager({'default': 'DENY'}, tmp_path))
    filesystem = register_filesystem_tools(denied)
    business = BusinessLedger(tmp_path / 'denied.db', tmp_path / 'denied', denied.permissions)
    register_product_workspace_tools(denied, store, filesystem, business)
    result = denied.execute('product_build_starter', {'product_id': product['id']})
    assert not result.success
    assert store.details(product['id'])['assets'] == []
