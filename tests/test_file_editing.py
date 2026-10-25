from pathlib import Path

import pytest

from app.tools.file_editing import FileEditing
from app.tools.filesystem import FilesystemTools
from app.safety.permissions import PermissionManager


@pytest.fixture
def editor(tmp_path):
    return FileEditing(FilesystemTools(PermissionManager({'default': 'ALLOW'}, tmp_path)))


def test_edits_preserve_bom_newlines_and_original_backup(editor, tmp_path):
    path = tmp_path / 'notes.txt'
    original = b'\xef\xbb\xbfhello\r\nworld\r\n'
    path.write_bytes(original)
    read = editor.read_editable_file(str(path))
    saved = editor.edit_file(str(path), 'world', 'reader', read['sha256'])
    assert path.read_bytes() == original.replace(b'world', b'reader')
    assert Path(saved['backup']).read_bytes() == original
    with pytest.raises(ValueError, match='changed'):
        editor.save_edited_file(str(path), 'stale content', read['sha256'])
    assert b'reader' in path.read_bytes()


def test_ambiguous_edits_and_binary_appends_do_not_change_files(editor, tmp_path):
    path = tmp_path / 'notes.txt'
    path.write_text('same same', encoding='utf-8')
    read = editor.read_editable_file(str(path))
    with pytest.raises(ValueError, match='Match count'):
        editor.edit_file(str(path), 'same', 'new', read['sha256'])
    assert path.read_text() == 'same same'
    binary = tmp_path / 'report.pdf'
    binary.write_bytes(b'%PDF fake')
    with pytest.raises(ValueError):
        editor.filesystem.append_file(str(binary), 'corruption')
    assert binary.read_bytes() == b'%PDF fake'


def test_native_generation_and_spreadsheet_edit(editor):
    import openpyxl
    from pptx import Presentation
    result = editor.write_spreadsheet('planner.xlsx', [['Task', 'Value'], ['Design', 2], ['Total', '=SUM(B2:B2)']])
    changed = editor.edit_spreadsheet_cells(result['path'], [{'sheet': 'Sheet1', 'cell': 'B2', 'value': 3}], result['sha256'])
    book = openpyxl.load_workbook(changed['path'])
    assert book.active['B2'].value == 3
    assert book.active['B3'].value == '=SUM(B2:B2)'
    book.close()
    slides = editor.write_presentation('kit.pptx', [{'title': 'Before', 'body': 'Details'}])
    deck = Presentation(slides['path'])
    assert deck.slides[0].shapes.title.text == 'Before'
    read = editor.read_document_for_edit(slides['path'])
    editor.edit_document_text(slides['path'], 'Before', 'After', read['sha256'])
    assert Presentation(slides['path']).slides[0].shapes.title.text == 'After'


def test_word_edits_cross_runs_without_losing_unrelated_formatting(editor, tmp_path):
    from docx import Document
    path = tmp_path / 'guide.docx'
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run('hello ').bold = True
    paragraph.add_run('world').italic = True
    paragraph.add_run('! untouched').underline = True
    document.save(path)
    read = editor.read_document_for_edit(str(path))
    result = editor.edit_document_text(str(path), 'hello world', 'welcome', read['sha256'])
    changed = Document(path).paragraphs[0]
    assert changed.text == 'welcome! untouched'
    assert changed.runs[0].bold and changed.runs[2].underline
    assert Document(result['backup']).paragraphs[0].text == 'hello world! untouched'


def test_protected_paths_stay_protected(editor, tmp_path):
    outside = tmp_path.parent / 'outside.txt'
    with pytest.raises(PermissionError):
        editor.read_editable_file(str(outside))
    with pytest.raises(ValueError):
        editor.filesystem.write_file('not-really.xlsx', 'raw text')


def test_spreadsheet_read_is_bounded_and_text_does_not_become_a_formula(editor):
    created = editor.write_spreadsheet('literal.xlsx', [['Header', '=1+1']])
    preview = editor.read_spreadsheet_for_edit(created['path'], cell_range='A1:B2')
    assert preview['cells'][1]['type'] == 'f'
    saved = editor.edit_spreadsheet_cells(created['path'], [{'sheet': 'Sheet1', 'cell': 'A1', 'value': '=literal text', 'value_type': 'text'}], preview['sha256'])
    checked = editor.read_spreadsheet_for_edit(saved['path'])
    assert checked['cells'][0]['type'] == 's' and checked['cells'][0]['value'] == '=literal text'
    assert checked['cells'][1]['value'] == '=1+1'
    with pytest.raises(ValueError, match='2000'):
        editor.read_spreadsheet_for_edit(saved['path'], cell_range='A1:ZZ1000')
    with pytest.raises(ValueError):
        editor.read_spreadsheet_for_edit(saved['path'], cell_range='A:A')


@pytest.mark.parametrize('value', [float('nan'), float('inf'), 'x' * 32768], ids=['nan', 'infinity', 'oversized-text'])
def test_invalid_excel_values_do_not_silently_lose_data(editor, value):
    created = editor.write_spreadsheet('checked.xlsx', [['Original']])
    before = Path(created['path']).read_bytes()
    with pytest.raises(ValueError):
        editor.edit_spreadsheet_cells(created['path'], [{'sheet': 'Sheet1', 'cell': 'A1', 'value': value}], created['sha256'])
    assert Path(created['path']).read_bytes() == before
