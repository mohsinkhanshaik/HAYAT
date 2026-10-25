import json
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

from app.agent.planner import Planner
from app.intent.router import IntentRouter, _business_command
from app.safety.permissions import PermissionManager
from app.tools.business import BusinessLedger, period_bounds, register_business_tools
from app.tools.products import ProductStore
from app.tools.registry import ToolRegistry


@pytest.fixture
def books(tmp_path):
    permissions = PermissionManager({"default": "ALLOW"}, tmp_path)
    products = ProductStore(tmp_path / "products.db", tmp_path / "bundles", permissions)
    ledger = BusinessLedger(tmp_path / "business.db", tmp_path / "business", permissions, products=products)
    registry = ToolRegistry(permissions)
    register_business_tools(registry, ledger)
    ledger.set_profile(business_name="Studio", payment_instructions="Bank transfer", default_currency="USD")
    ledger.add_client("Acme Ltd", email="ap@acme.test", hourly_rate="40")
    return ledger, registry, products


def pdf_text(path):
    import pdfplumber

    with pdfplumber.open(path) as document:
        return "\n".join(page.extract_text() or "" for page in document.pages)


def test_quote_to_invoice_with_logged_time_tax_and_pdf(books):
    ledger, _, _ = books
    ledger.log_time("acme", "2.5", "Landing page")
    quote = ledger.create_quote("Acme Ltd", [{"description": "Logo", "quantity": 1, "unit_price": "150"}], tax_percent="5")
    assert quote["number"] == f"Q-{date.today().year}-0001" and quote["total"] == "157.50 USD"
    invoice = ledger.create_invoice("Acme Ltd", include_unbilled_time=True, from_quote=quote["number"])
    assert invoice["number"] == f"INV-{date.today().year}-0001"
    assert invoice["subtotal_cents"] == 25000 and invoice["tax_cents"] == 1250 and invoice["total"] == "262.50 USD"
    text = pdf_text(invoice["path"])
    for expected in ("INVOICE", invoice["number"], "Acme Ltd", "Logo", "Landing page", "262.50 USD", "Bank transfer"):
        assert expected in text
    # Billed time and an invoiced quote cannot be billed twice.
    assert ledger.unbilled_time()["unbilled"] == []
    with pytest.raises(ValueError, match="already"):
        ledger.create_invoice("Acme Ltd", from_quote=quote["number"])
    assert ledger.create_invoice("Acme Ltd", [{"description": "x", "unit_price": "1"}])["number"].endswith("0002")


def test_payments_partial_duplicate_overpay_and_status(books):
    ledger, _, _ = books
    invoice = ledger.create_invoice("Acme Ltd", [{"description": "Work", "quantity": "2", "unit_price": "50"}])
    first = ledger.record_payment(invoice["number"], "40", reference="bank-1")
    assert first["recorded"] and first["status"] == "unpaid" and first["balance"] == "60.00 USD"
    assert ledger.record_payment(invoice["number"], "40", reference="bank-1")["duplicate"]
    with pytest.raises(ValueError, match="more than"):
        ledger.record_payment(invoice["number"], "61")
    assert ledger.record_payment(invoice["number"])["status"] == "paid"
    with pytest.raises(ValueError, match="Nothing is owed"):
        ledger.record_payment(invoice["number"])
    with pytest.raises(ValueError, match="can't be voided"):
        ledger.void_invoice(invoice["number"], "mistake")


def test_void_returns_time_to_unbilled(books):
    ledger, _, _ = books
    ledger.log_time("Acme Ltd", "1", "Fix")
    invoice = ledger.create_invoice("Acme Ltd", include_unbilled_time=True)
    assert ledger.void_invoice(invoice["number"], "wrong client")["status"] == "void"
    assert ledger.unbilled_time()["unbilled"][0]["hours"] == 1
    assert ledger.invoices("unpaid")["count"] == 0


def test_overdue_and_report_combines_clients_products_expenses_and_goal(books):
    ledger, _, products = books
    invoice = ledger.create_invoice("Acme Ltd", [{"description": "Old work", "unit_price": "300"}], due_days=0)
    with ledger._db() as db:
        db.execute("UPDATE documents SET due=? WHERE number=?", ((date.today() - timedelta(days=3)).isoformat(), invoice["number"]))
    paid = ledger.create_invoice("Acme Ltd", [{"description": "New work", "unit_price": "100"}])
    ledger.record_payment(paid["number"])
    product = products.create("Planner", "Freelancers", "Template", "10", "Personal use")
    products.record_sale(product["id"], "order-1", "10", "1")
    ledger.record_expense("12.50", "software", "Domain")
    ledger.set_goal("1000")
    ledger.add_lead("Theme gig", "Upwork", "https://example.com/job", "500", follow_up="today")
    report = ledger.report("this month")
    usd = report["by_currency"][0]
    assert (usd["client_income"], usd["product_sales_net"], usd["expenses"], usd["profit_before_tax"]) == \
        ("100.00 USD", "9.00 USD", "12.50 USD", "96.50 USD")
    assert [d["number"] for d in report["overdue"]] == [invoice["number"]]
    assert report["goal"]["earned"] == "109.00 USD" and report["goal"]["percent"] == 10.9
    assert any("overdue" in action for action in report["next_actions"])
    assert any("Follow up" in action for action in report["next_actions"])
    assert "not verified" in report["basis"]
    assert ledger.report("last year")["by_currency"] == []


def test_currencies_never_mix(books):
    ledger, _, _ = books
    ledger.add_client("Euro Co", hourly_rate="50", currency="EUR")
    eur = ledger.create_invoice("Euro Co", [{"description": "x", "unit_price": "10"}])
    ledger.record_payment(eur["number"])
    usd = ledger.create_invoice("Acme Ltd", [{"description": "x", "unit_price": "10"}])
    ledger.record_payment(usd["number"])
    assert {row["currency"] for row in ledger.report()["by_currency"]} == {"EUR", "USD"}
    ledger.log_time("Euro Co", "1", "work")
    with pytest.raises(ValueError, match="client's currency"):
        ledger.create_invoice("Euro Co", include_unbilled_time=True, currency="USD")
    with pytest.raises(ValueError, match="currencies"):
        ledger.add_client("Yen Co", currency="JPY")


def test_timer_leads_and_validation(books):
    ledger, _, _ = books
    started = ledger.start_timer("Acme Ltd", "Call")
    with pytest.raises(ValueError, match="already running"):
        ledger.start_timer("Acme Ltd", "Second")
    with ledger._db() as db:
        db.execute("UPDATE time_entries SET running_since=? WHERE running_since=?",
                   ("2000-01-01T00:00:00+00:00", started["started"]))
    assert ledger.stop_timer()["hours"] == 24  # capped, never days of runaway time
    with pytest.raises(ValueError, match="No timer"):
        ledger.stop_timer()
    lead = ledger.add_lead("Site rebuild", "Referral", value="900")
    assert ledger.update_lead(lead["lead_id"], status="won")["status"] == "won"
    assert ledger.leads("open")["count"] == 0 and ledger.leads("won")["count"] == 1
    with pytest.raises(ValueError):
        ledger.update_lead(lead["lead_id"], status="maybe")
    with pytest.raises(ValueError, match="http"):
        ledger.add_lead("x", "y", url="javascript:alert(1)")
    with pytest.raises(ValueError, match="Unknown client"):
        ledger.log_time("Nobody", "1", "x")
    with pytest.raises(ValueError):
        ledger.log_time("Acme Ltd", "25", "x")
    with pytest.raises(ValueError):
        period_bounds("next decade")


def test_export_workbook_has_every_sheet(books):
    ledger, _, _ = books
    invoice = ledger.create_invoice("Acme Ltd", [{"description": "Work", "unit_price": "100"}])
    ledger.record_payment(invoice["number"], reference="r1")
    ledger.record_expense("20", "ads", "Promo")
    result = ledger.export()
    book = load_workbook(result["path"])
    assert book.sheetnames == ["Income received", "Expenses", "Invoices issued", "Other income", "Withdrawals",
                               "Product sales", "Read me"]
    assert book["Income received"].cell(2, 5).value == 100
    assert book["Expenses"].cell(2, 5).value == 20


def test_tools_registered_with_permissions(books, tmp_path):
    ledger, registry, _ = books
    result = registry.execute("invoice_create", {"client": "Acme Ltd", "items": [
        {"description": "Design", "quantity": "3", "unit_price": 40}]})
    assert result.success, result.error
    assert registry.execute("business_report", {}).success
    json.dumps(registry.execute("invoice_list", {"status": "unpaid"}).output)
    denied = ToolRegistry(PermissionManager({"default": "DENY"}, tmp_path))
    register_business_tools(denied, ledger)
    assert not denied.execute("expense_record", {"amount": "5", "category": "x", "description": "y"}).success
    assert ledger.expenses("all")["expenses"] == []


def test_files_stay_inside_allowed_roots(tmp_path):
    permissions = PermissionManager({"default": "ALLOW"}, tmp_path / "project")
    ledger = BusinessLedger(tmp_path / "project" / "b.db", tmp_path / "outside", permissions)
    ledger.add_client("Acme")
    with pytest.raises(PermissionError):
        ledger.create_invoice("Acme", [{"description": "x", "unit_price": "1"}])
    # No PDF means no invoice: the books and the files stay consistent.
    assert ledger.invoices()["count"] == 0


@pytest.mark.parametrize("text", [
    "Create an invoice for Acme for 3 hours of design at 40", "log 2.5 hours for Acme on the landing page",
    "Show my unpaid invoices", "How much did I earn this month?", "Add a lead: logo gig on Upwork",
    "Record an expense of 12 dollars for software", "Set my income goal to 2000 dollars",
    "Start a work timer for Acme", "export my business records for taxes", "add client Acme with rate 40",
])
def test_business_commands_reach_the_executor(text):
    assert _business_command(text)
    assert IntentRouter().route(text).requires_tools


@pytest.mark.parametrize("text", [
    "What is an invoice?", "how do I write a quote", "set a timer for 5 minutes",
    "send a whatsapp message to my client Ali saying hi", "start chrome and show my clients website",
    "check my email for client messages", "update the client app",
])
def test_other_requests_keep_their_routes(text):
    assert not _business_command(text)


@pytest.mark.parametrize("goal, tool, arguments", [
    ("Show my business report", "business_report", {}),
    ("How much did I earn last month?", "business_report", {"period": "last month"}),
    ("Show my overdue invoices", "invoice_list", {"status": "overdue"}),
    ("show my clients", "client_list", {}),
    ("show follow-ups due", "lead_list", {"status": "open", "due_only": True}),
    ("stop the work timer", "time_timer_stop", {}),
    ("show my expenses for this year", "expense_list", {"period": "this year"}),
    ("export my business records for taxes for 2026", "business_export", {"year": "2026"}),
])
def test_read_requests_plan_directly(books, goal, tool, arguments):
    _, registry, _ = books
    plan = Planner(registry).create_plan(goal)
    assert (plan[0].tool, plan[0].arguments) == (tool, arguments)


def test_write_requests_offer_the_business_tools(books):
    _, registry, _ = books
    candidates = Planner(registry)._candidate_tools("Create an invoice for Acme for 3 hours of design")
    assert {"invoice_create", "time_log", "client_add"} <= set(candidates)
    finder = Planner(registry)._candidate_tools("Find freelance gigs I could do and track them as leads")
    assert "lead_add" in finder


def test_wallet_collects_every_kind_of_earning_and_subtracts_costs(books):
    ledger, _, products = books
    invoice = ledger.create_invoice("Acme Ltd", [{"description": "Work", "unit_price": "200"}])
    ledger.record_payment(invoice["number"])
    product = products.create("Planner", "Freelancers", "Template", "10", "Personal use")
    products.record_sale(product["id"], "order-1", "10", "1")
    assert ledger.record_income("50", "Fiverr", "Logo gig payout", reference="fv-1")["recorded"]
    assert ledger.record_income("50", "Fiverr", "Logo gig payout", reference="fv-1")["duplicate"]
    ledger.record_expense("20", "software", "Domain")
    wallet = ledger.wallet()
    usd = wallet["balances"][0]
    assert (usd["earned"], usd["expenses"], usd["available"]) == ("259.00 USD", "20.00 USD", "239.00 USD")
    kinds = {item["kind"] for item in wallet["activity"]}
    assert {"Invoice payment", "Product sale", "Income", "Expense"} <= kinds
    assert "not the money" in wallet["basis"]
    # Other income counts as earnings in the report, too.
    assert ledger.report()["by_currency"][0]["other_income"] == "50.00 USD"


def test_withdrawals_wait_for_a_method_and_never_exceed_the_balance(books):
    ledger, _, _ = books
    ledger.record_income("100", "Upwork", "Payout")
    report = ledger.report()
    assert any("withdrawal method" in action for action in report["next_actions"])
    with pytest.raises(ValueError, match="Set up a withdrawal method"):
        ledger.withdraw("10")
    ledger.set_payout("paypal", "me@example.com")
    pending = ledger.withdraw("60")
    assert pending["status"] == "pending" and pending["available_after"] == "40.00 USD"
    assert "does not move money" in pending["next_step"] and "PayPal" in pending["next_step"]
    with pytest.raises(ValueError, match="Only 40.00 USD"):
        ledger.withdraw("41")
    assert ledger.complete_withdrawal(pending["withdrawal_id"], "PP-1")["status"] == "done"
    with pytest.raises(ValueError, match="already done"):
        ledger.complete_withdrawal(pending["withdrawal_id"])
    rest = ledger.withdraw("all")
    assert rest["amount"] == "40.00 USD"
    ledger.cancel_withdrawal(rest["withdrawal_id"])
    usd = ledger.wallet(activity=0)["balances"][0]
    assert (usd["withdrawn"], usd["pending"], usd["available"]) == ("60.00 USD", "0.00 USD", "40.00 USD")


@pytest.mark.parametrize("method, name, last4", [
    ("bank", "Account 12345678", ""), ("bank", "Card 4111-1111-1111-1111", ""),
    ("paypal", "password hunter2", ""), ("bank", "Savings", "12345"), ("bitcoin", "Wallet", ""),
])
def test_payout_setup_refuses_numbers_secrets_and_unknown_methods(books, method, name, last4):
    ledger, _, _ = books
    with pytest.raises(ValueError):
        ledger.set_payout(method, name, last4)
    assert not ledger.payout()["configured"]


def test_payout_keeps_only_a_name_and_last_four(books):
    ledger, _, _ = books
    saved = ledger.set_payout("bank", "My savings", "4321", "USD")
    assert saved["destination"] == "Bank: My savings ending 4321"
    assert not ledger.clear_payout()["configured"]


@pytest.mark.parametrize("goal, tool, arguments", [
    ("show my wallet", "wallet_balance", {}),
    ("What's my balance?", "wallet_balance", {}),
    ("How much can I withdraw?", "wallet_balance", {}),
    ("withdraw $50 from my wallet", "wallet_withdraw", {"amount": "50"}),
    ("withdraw everything", "wallet_withdraw", {"amount": "everything"}),
    ("withdrawal 3 is done", "wallet_withdrawal_done", {"withdrawal_id": 3}),
    ("cancel withdrawal 2", "wallet_withdrawal_cancel", {"withdrawal_id": 2}),
    ("set up withdrawal to PayPal", "wallet_payout_setup", {"method": "paypal", "name": "My PayPal account"}),
    ("set up my payout to bank called my savings ending 4321", "wallet_payout_setup",
     {"method": "bank", "name": "my savings", "last4": "4321"}),
])
def test_wallet_requests_plan_directly(books, goal, tool, arguments):
    _, registry, _ = books
    assert _business_command(goal)
    plan = Planner(registry).create_plan(goal)
    assert (plan[0].tool, plan[0].arguments) == (tool, arguments)


@pytest.mark.parametrize("text", ["what is a wallet", "set the audio balance", "how much does a wallet cost",
                                  "cancel the meeting"])
def test_wallet_words_do_not_capture_other_requests(text):
    assert not _business_command(text)
