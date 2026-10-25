from __future__ import annotations

import json
import sys
import threading

import httpx
import pytest

from app.attachments import AttachmentReader
from app.automation.event_engine import EventRule, EventTriggerEngine
from app.events import EventBus
from app.integrations.browser_bridge import BrowserExtensionBridge
from app.integrations.connectors import ConnectorGateway, ConnectorManifest, ConnectorRegistry
from app.integrations.mcp_gateway import MCPGateway, MCPServerManifest, register_builtin_web_research
from app.integrations.secrets import SecretBroker


class MemorySecrets:
    def __init__(self):
        self.values = {}

    def set_password(self, service, username, password):
        self.values[(service, username)] = password

    def get_password(self, service, username):
        return self.values.get((service, username))

    def delete_password(self, service, username):
        self.values.pop((service, username), None)


def broker(tmp_path, bus=None):
    return SecretBroker(tmp_path / "secrets.db", backend=MemorySecrets(), event_bus=bus)


def test_secret_broker_uses_scoped_single_use_leases_and_never_lists_values(tmp_path):
    item = broker(tmp_path)
    item.store("github-token", "never-show-this-value", scopes=["connector:github"])
    lease = item.issue_lease("github-token", "connector:github")
    assert item.redeem(lease, "connector:github") == "never-show-this-value"
    with pytest.raises(PermissionError):
        item.redeem(lease, "connector:github")
    with pytest.raises(PermissionError):
        item.issue_lease("github-token", "connector:notion")
    assert "never-show-this-value" not in json.dumps(item.status())


def test_connector_gateway_separates_read_and_write_and_injects_secret_without_returning_it(tmp_path, monkeypatch):
    monkeypatch.setattr("app.integrations.connectors.validate_public_url", lambda value: value)
    item = broker(tmp_path)
    item.store("github-token", "connector-secret-value", scopes=["connector:github"])
    registry = ConnectorRegistry(tmp_path / "connectors.db", catalog=tmp_path / "none.yaml")
    registry.install(ConnectorManifest("github", "GitHub", "http", "https://api.example.com",
                                       ["repo:read"], ["repos"], "github-token"), enable=True)
    seen = {}

    def request(method, url, **kwargs):
        seen.update(method=method, url=url, **kwargs)
        return httpx.Response(200, json={"repositories": 3}, request=httpx.Request(method, url))

    gateway = ConnectorGateway(registry, item, request=request)
    result = gateway.call("github", "repos")
    assert result["output"] == {"repositories": 3}
    assert seen["headers"]["Authorization"] == "Bearer connector-secret-value"
    assert "connector-secret-value" not in json.dumps(result)
    with pytest.raises(PermissionError):
        gateway.call("github", "issues", write=False)


def test_mcp_registry_requires_installed_executable_and_redacts_secret_names(tmp_path):
    gateway = MCPGateway(tmp_path / "mcp.db", secrets=broker(tmp_path))
    result = gateway.register(MCPServerManifest(
        "local-test", "Local Test", sys.executable, ["-m", "example"], ["read"],
        {"API_TOKEN": "internal-secret-name"}), enable=False)
    assert result["enabled"] is False
    assert result["secret_environment"] == ["API_TOKEN"]
    assert "internal-secret-name" not in json.dumps(result)
    with pytest.raises(FileNotFoundError):
        gateway.register(MCPServerManifest("missing", "Missing", "definitely-not-installed"))


def test_builtin_zero_cost_web_mcp_is_enabled_and_reports_google_boundary(tmp_path):
    gateway = MCPGateway(tmp_path / "mcp.db", secrets=broker(tmp_path))
    manifest = register_builtin_web_research(gateway)
    assert manifest["enabled"] is True
    assert "deep_web_research" in manifest["read_only_tools"]
    names = {tool["name"] for tool in gateway.tools(manifest["server_id"])}
    assert {"search_web", "deep_web_research", "fetch_public_page",
            "inspect_google_snapshot", "google_search_status"} <= names
    result = gateway.call(manifest["server_id"], "google_search_status")
    serialized = json.dumps(result)
    assert "direct_results_page_automation" in serialized
    assert "false" in serialized.lower()


def test_google_snapshot_inspection_requires_an_explicit_google_page():
    from app.integrations.web_research_mcp import inspect_google_snapshot

    result = inspect_google_snapshot(
        "Example result https://example.com/source", "https://www.google.com/search?q=example",
    )
    assert result["live_google_automation"] is False
    assert result["links"] == ["https://example.com/source"]
    with pytest.raises(ValueError):
        inspect_google_snapshot("content", "https://example.com/")


def test_browser_bridge_authenticates_and_returns_a_live_extension_snapshot(tmp_path):
    bridge = BrowserExtensionBridge(broker(tmp_path), port=18765)
    bridge.start()
    try:
        from websockets.sync.client import connect

        with connect("ws://127.0.0.1:18765") as socket:
            socket.send(json.dumps({"type": "authenticate", "token": bridge.pairing_token()}))
            assert json.loads(socket.recv())["type"] == "authenticated"

            def extension():
                command = json.loads(socket.recv())
                assert command["command"] == "snapshot"
                socket.send(json.dumps({"type": "response", "id": command["id"], "ok": True,
                                        "result": {"url": "https://example.com/", "title": "Example",
                                                   "text": "A live page", "links": [], "forms": []}}))

            worker = threading.Thread(target=extension)
            worker.start()
            result = bridge.snapshot()
            worker.join(timeout=2)
            assert result["text"] == "A live page"
            assert bridge.status()["latest_page"]["title"] == "Example"
    finally:
        bridge.stop()


def test_event_rules_trigger_suggestions_and_enforce_cooldown(tmp_path):
    bus = EventBus(tmp_path / "events.db", durable=False)
    suggestions = []
    engine = EventTriggerEngine(bus, path=tmp_path / "rules.db",
                                suggest=lambda goal, event: suggestions.append((goal, event["event_type"])))
    engine.add(EventRule("RESOURCE_ALERT", "Inspect resource pressure", {"severity": "high"},
                         cooldown_seconds=60, enabled=True))
    engine.start()
    try:
        bus.publish("RESOURCE_ALERT", "test", {"severity": "low"})
        bus.publish("RESOURCE_ALERT", "test", {"severity": "high"})
        bus.publish("RESOURCE_ALERT", "test", {"severity": "high"})
        assert suggestions == [("Inspect resource pressure", "RESOURCE_ALERT")]
        assert engine.list()[0]["run_count"] == 1
    finally:
        engine.stop(); bus.close()


def test_document_intelligence_extracts_word_excel_and_powerpoint(tmp_path):
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation

    word_path = tmp_path / "brief.docx"
    word = Document(); word.add_heading("Project Atlas", 1); word.add_paragraph("Budget is $42,000.")
    table = word.add_table(rows=1, cols=2); table.rows[0].cells[0].text = "Owner"; table.rows[0].cells[1].text = "Mohsin"
    word.save(word_path)

    excel_path = tmp_path / "metrics.xlsx"
    book = Workbook(); sheet = book.active; sheet.title = "KPIs"; sheet.append(["Metric", "Value"]); sheet.append(["Availability", "99.99%"]); book.save(excel_path)

    slides_path = tmp_path / "review.pptx"
    slides = Presentation(); slide = slides.slides.add_slide(slides.slide_layouts[5]); slide.shapes.title.text = "Quarterly Review"; slides.save(slides_path)

    reader = AttachmentReader()
    assert "Project Atlas" in reader.read(word_path).text
    assert "Availability | 99.99%" in reader.read(excel_path).text
    assert "Quarterly Review" in reader.read(slides_path).text
