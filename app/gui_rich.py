"""Making the chat readable and the work visible.

Four things the owner asked for on 2026-10-04 ("clear, understandable,
presentable; show image or doc files in chat and open them; show the working
while computer use is going; whatever runs in the background is shown, nothing
hidden"):

* replies arrive as Markdown with LaTeX maths -- "$\\$120 \\times 0.8 = \\$96$" --
  which a Tk text widget showed as raw symbols; :func:`clean_math` and
  :func:`markdown_segments` turn them into plain, styled text;
* images and documents named in a message are found (:func:`find_images`,
  :func:`find_documents`) and previewed (:func:`document_preview`);
* every running operation gets a plain-language line (:func:`activity_line`);
* :class:`WorkingView` shows each step of a turn as it happens, with a
  screenshot after each step that acts on the computer.

Kept free of the main window so each piece can be tested without Tk.
"""
from __future__ import annotations

import re
import time
from app.gui_resources import StringVar

import tkinter as tk
from dataclasses import dataclass
from pathlib import Path
from typing import Any

# ----------------------------------------------------------------- maths

_LATEX_SYMBOLS = {
    r"\times": "×", r"\cdot": "·", r"\div": "÷", r"\pm": "±", r"\le": "≤", r"\leq": "≤", r"\ge": "≥",
    r"\geq": "≥", r"\ne": "≠", r"\neq": "≠", r"\approx": "≈", r"\infty": "∞", r"\rightarrow": "→",
    r"\to": "→", r"\Rightarrow": "⇒", r"\leftarrow": "←", r"\sum": "Σ", r"\pi": "π", r"\alpha": "α",
    r"\beta": "β", r"\theta": "θ", r"\Delta": "Δ", r"\delta": "δ", r"\mu": "μ", r"\sigma": "σ",
    r"\percent": "%", r"\quad": " ", r"\qquad": "  ", r"\,": " ", r"\;": " ", r"\!": "", r"\ ": " ",
    r"\left": "", r"\right": "",
}
_SUPERSCRIPT = str.maketrans("0123456789+-()n", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁽⁾ⁿ")


def clean_math(text: str) -> str:
    """LaTeX maths as plain text a person can read: "$\\$80 \\times 1.5 = \\$120$" -> "$80 × 1.5 = $120".

    Only text inside maths markers is converted, so a Windows path such as
    D:\\tools is never mistaken for the command \\to, and "$5 and $10" stays money.
    """
    value = str(text or "")
    if "$" not in value and "\\[" not in value and "\\(" not in value:
        return value
    value = re.sub(r"\$\$(.+?)\$\$", lambda m: _convert_math(m.group(1)), value, flags=re.S)
    value = re.sub(r"\\\[(.+?)\\\]", lambda m: _convert_math(m.group(1)), value, flags=re.S)
    value = re.sub(r"\\\((.+?)\\\)", lambda m: _convert_math(m.group(1)), value, flags=re.S)
    value = re.sub(r"(?<!\\)\$((?:\\\$|[^$\n])+?)(?<!\\)\$",
                   lambda m: _convert_math(m.group(1)) if re.search(r"\\[A-Za-z$%]|\^|_\{|=", m.group(1))
                   else m.group(0), value)
    return value


def _convert_math(inner: str) -> str:
    value = inner
    # \text{...}, \mathrm{...}, \textbf{...} keep their words.
    value = re.sub(r"\\(?:text|mathrm|mathbf|textbf|operatorname|boxed)\{([^{}]*)\}", r"\1", value)
    for _ in range(3):  # nested fractions, innermost first
        value = re.sub(r"\\[dt]?frac\{([^{}]*)\}\{([^{}]*)\}", r"(\1)/(\2)", value)
    value = re.sub(r"\\sqrt\{([^{}]*)\}", r"√(\1)", value)
    value = re.sub(r"\^\{([0-9+\-()n]+)\}", lambda m: m.group(1).translate(_SUPERSCRIPT), value)
    value = re.sub(r"\^([0-9n])", lambda m: m.group(1).translate(_SUPERSCRIPT), value)
    for command in sorted(_LATEX_SYMBOLS, key=len, reverse=True):
        value = re.sub(re.escape(command) + r"(?![A-Za-z])", lambda _m, s=_LATEX_SYMBOLS[command]: s, value)
    value = value.replace(r"\$", "$").replace(r"\%", "%").replace(r"\&", "&").replace(r"\_", "_")
    value = re.sub(r"\{([^{}]*)\}", r"\1", value)  # leftover grouping braces
    return re.sub(r"[ \t]{2,}", " ", value).strip()


# ----------------------------------------------------------------- markdown

@dataclass(frozen=True)
class Segment:
    text: str
    tags: tuple[str, ...] = ()
    href: str = ""


_INLINE = re.compile(r"(\[[^\]\n]+\]\(https?://[^\s)]+\)|\*\*[^*\n]+?\*\*|__[^_\n]+?__|`[^`\n]+?`|\*[^*\s][^*\n]*?\*)")


def _inline(line: str, base: tuple[str, ...]) -> list[Segment]:
    parts: list[Segment] = []
    for piece in _INLINE.split(line):
        if not piece:
            continue
        link = re.fullmatch(r"\[([^\]\n]+)\]\((https?://[^\s)]+)\)", piece)
        if link:
            parts.append(Segment(link.group(1), base + ("md_link",), link.group(2)))
        elif (piece.startswith("**") and piece.endswith("**")) or (piece.startswith("__") and piece.endswith("__")):
            parts.append(Segment(piece[2:-2], base + ("md_bold",)))
        elif piece.startswith("`") and piece.endswith("`"):
            parts.append(Segment(piece[1:-1], base + ("md_code",)))
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            parts.append(Segment(piece[1:-1], base + ("md_italic",)))
        else:
            parts.append(Segment(piece, base))
    return parts


def markdown_segments(text: str) -> list[Segment]:
    """A reply as styled pieces: headings, bullets, numbered steps, bold, code and code blocks."""
    segments: list[Segment] = []
    in_code = False
    for raw_line in clean_math(text).split("\n"):
        line = raw_line.rstrip()
        if line.strip().startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            segments.append(Segment(line + "\n", ("md_codeblock",)))
            continue
        heading = re.match(r"^\s{0,3}(#{1,4})\s+(.*)$", line)
        bullet = re.match(r"^(\s*)[-*•]\s+(.*)$", line)
        numbered = re.match(r"^(\s*)(\d{1,3}[.)])\s+(.*)$", line)
        if heading:
            segments.extend(_inline(heading.group(2).strip("# "), ("md_heading",)))
        elif re.fullmatch(r"\s*(?:-{3,}|\*{3,}|_{3,})\s*", line):
            segments.append(Segment("─" * 24, ("md_rule",)))
        elif bullet:
            level = "md_bullet2" if len(bullet.group(1)) >= 2 else "md_bullet"
            segments.append(Segment("• ", (level,)))
            segments.extend(_inline(bullet.group(2), (level,)))
        elif numbered:
            level = "md_bullet2" if len(numbered.group(1)) >= 2 else "md_bullet"
            segments.append(Segment(numbered.group(2) + " ", (level, "md_bold")))
            segments.extend(_inline(numbered.group(3), (level,)))
        else:
            segments.extend(_inline(line, ()))
        segments.append(Segment("\n"))
    while segments and segments[-1].text == "\n":
        segments.pop()
    return segments


# ----------------------------------------------------------------- media

IMAGE_TYPES = (".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif")
DOCUMENT_TYPES = (".pdf", ".docx", ".txt", ".md", ".csv", ".json", ".xlsx", ".pptx", ".log", ".py", ".yaml", ".yml", ".html", ".css", ".js", ".ts", ".tsx", ".jsx", ".sql", ".toml", ".ini", ".zip")
_PATH = re.compile(r"(?i)(?:\"([a-z]:[\\/][^\"\r\n]+?)\"|([a-z]:[\\/][^\r\n\"<>|*?]+?\.(?:"
                   r"png|jpe?g|webp|bmp|gif|pdf|docx|txt|md|csv|json|xlsx|pptx|log|py|ya?ml|html|css|js|tsx?|jsx|sql|toml|ini|zip|mp4|mkv|mov|webm|m4v|avi))(?=$|[\s,;:!?)\]'\"]|\.\s|\.$))")


def _paths(text: str) -> list[Path]:
    found: list[Path] = []
    for match in _PATH.finditer(str(text or "")):
        candidate = (match.group(1) or match.group(2) or "").strip().rstrip(".,;:!?")
        try:
            path = Path(candidate)
            if path.is_file() and path not in found:
                found.append(path)
        except OSError:
            continue
    return found


def find_images(text: str, limit: int = 4) -> list[Path]:
    """Existing image files named in a message, in order, at most ``limit``."""
    return [path for path in _paths(text) if path.suffix.lower() in IMAGE_TYPES][:limit]


def find_documents(text: str, limit: int = 3) -> list[Path]:
    """Existing document files named in a message, in order, at most ``limit``."""
    return [path for path in _paths(text) if path.suffix.lower() in DOCUMENT_TYPES][:limit]


def find_videos(text: str, limit: int = 4) -> list[Path]:
    from app.video import VIDEO_KINDS
    return [path for path in _paths(text) if path.suffix.lower() in VIDEO_KINDS][:limit]


def find_video_links(text: str, limit: int = 3) -> list[str]:
    """Recognize video destinations without fetching or interpreting page content."""
    from urllib.parse import urlsplit
    platforms = {'youtube.com', 'youtu.be', 'vimeo.com', 'dailymotion.com', 'dai.ly',
                 'twitch.tv', 'tiktok.com', 'instagram.com', 'facebook.com', 'fb.watch'}
    found = []
    for match in re.finditer(r'https?://[^\s<>"\']+', text):
        url = match.group().rstrip('.,;!)')
        parsed = urlsplit(url)
        host = (parsed.hostname or '').lower().removeprefix('www.')
        if parsed.username or parsed.password:
            continue
        if host in platforms or any(host.endswith('.' + p) for p in platforms) or Path(parsed.path).suffix.lower() in {'.mp4', '.webm', '.mkv', '.mov', '.m4v'}:
            if url not in found:
                found.append(url)
        if len(found) == limit:
            break
    return found


def _size(path: Path) -> str:
    size = path.stat().st_size
    for unit in ("bytes", "KB", "MB", "GB"):
        if size < 1024 or unit == "GB":
            return f"{size:.0f} {unit}" if unit == "bytes" else f"{size:.1f} {unit}"
        size /= 1024
    return ""


def document_preview(path: Path, lines: int = 6) -> dict[str, Any]:
    """What to show on a document's card: its kind, a few lines, and for a PDF a first-page picture.

    Best-effort: any file that cannot be read still gets a card with its name
    and size, so it can always be opened.
    """
    suffix = path.suffix.lower()
    info: dict[str, Any] = {"name": path.name, "size": "unavailable", "kind": suffix.lstrip(".").upper(),
                            "lines": [], "thumbnail": None, "pages": None}
    try:
        info["size"] = _size(path)
        if suffix == ".zip":
            import zipfile

            # List members without extracting, executing or reading their contents.
            with zipfile.ZipFile(path) as archive:
                members = archive.infolist()
                info["lines"] = [f"{len(members)} files in bundle"] + [
                    f"{member.filename[:140]} ({member.file_size} bytes)" for member in members[:lines]]
        elif suffix in set(DOCUMENT_TYPES) - {".pdf", ".docx", ".xlsx", ".pptx", ".zip"}:
            with path.open(encoding="utf-8", errors="replace") as handle:
                info["lines"] = [handle.readline(4096).rstrip()[:160] for _ in range(lines)]
        elif suffix == ".pdf":
            from app.pdf_preview_worker import render_pdf

            rendered = render_pdf(path, thumbnail=True, lines=lines)
            info.update(pages=rendered['pages'], thumbnail=rendered['image'], lines=rendered['lines'])
        elif suffix == ".docx":
            import docx

            document = docx.Document(str(path))
            info["lines"] = [p.text.strip()[:160] for p in document.paragraphs if p.text.strip()][:lines]
        elif suffix == ".xlsx":
            import openpyxl

            book = openpyxl.load_workbook(path, read_only=True, data_only=True)
            sheet = book.worksheets[0]
            info["pages"] = len(book.worksheets)
            for row in sheet.iter_rows(max_row=lines, values_only=True):
                info["lines"].append(" | ".join("" if cell is None else str(cell) for cell in row)[:160])
            book.close()
        elif suffix == ".pptx":
            import pptx

            deck = pptx.Presentation(str(path))
            info["pages"] = len(deck.slides)
            for slide in deck.slides:
                titles = [shape.text_frame.text.strip() for shape in slide.shapes
                          if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip()]
                if titles:
                    info["lines"].append(titles[0][:160])
                if len(info["lines"]) >= lines:
                    break
    except Exception:
        info["lines"] = info["lines"] or []
    info["lines"] = [line for line in info["lines"] if line]
    return info


# ----------------------------------------------------------------- activity

_TOOL_WORDS = {
    "inspect_video_link": "Checking video qualities", "download_video": "Downloading a video",
    "analyze_video_clips": "Assessing clip moments", "create_video_clip": "Exporting a video clip",
    "make_video_clips": "Making video clips", "video_preview": "Preparing a chat video preview",
    "play_video": "Opening the chat video player", "assess_video": "Understanding a linked video",
    "video_player_control": "Controlling the chat player",
    "edit_video_clip": "Editing a video", "share_video_result": "Creating a local video result link",
    "computer_click": "Clicking", "computer_type": "Typing", "computer_press": "Pressing keys",
    "computer_scroll": "Scrolling", "computer_screenshot": "Taking a screenshot", "computer_move": "Moving the mouse",
    "computer_move_mouse": "Moving the cursor", "computer_mouse_position": "Checking cursor position",
    "computer_drag": "Dragging", "computer_status": "Checking desktop input",
    "launch_application": "Opening an app", "close_application": "Closing an app",
    "browser_launch": "Starting the browser", "browser_navigate": "Opening a web page",
    "browser_click": "Clicking on the page", "browser_type": "Typing on the page",
    "browser_read_page": "Reading the page", "browser_screenshot": "Capturing the page",
    "web_search": "Searching the web", "research_topic": "Researching", "public_web_fetch": "Reading a web page",
    "read_file": "Reading a file", "write_file": "Writing a file", "list_folder": "Looking in a folder",
    "describe_image": "Looking at the image", "describe_screen": "Looking at the screen",
    "generate_image": "Making an image", "generate_video": "Making a video", "knowledge_search": "Searching memory",
    "camera_observe": "Looking through the camera",
}
_KIND_WORDS = {"model": "Thinking", "learning": "Background learning", "research": "Researching",
               "image": "Making an image", "task": "Working on a task", "voice": "Voice",
               "service": "Model service"}
_KIND_ICONS = {"tool": "⚙", "model": "✦", "learning": "📚", "research": "🔎", "image": "🖼",
               "task": "▶", "voice": "🎙", "service": "⛁"}
#: Tools that act on the computer, after which a screenshot shows the result.
COMPUTER_ACTIONS = ("computer_", "launch_application", "close_application", "browser_click", "browser_type",
                    "browser_navigate")


def activity_line(item: dict[str, Any]) -> str:
    """One running operation in plain words: "⚙ Opening a web page: example.com  4s"."""
    kind = str(item.get("kind") or "")
    label = str(item.get("label") or "")
    detail = " ".join(str(item.get("detail") or "").split())
    if kind == "tool":
        words = _TOOL_WORDS.get(label) or label.replace("_", " ").capitalize()
    elif kind == "model":
        words = "Thinking" + (f" ({detail.split('/')[-1]})" if detail else "")
        detail = ""
    else:
        words = _KIND_WORDS.get(kind, kind.capitalize() or "Working") + (f" · {label}" if kind == "service" else "")
    if detail:
        words += f": {detail[:90]}"
    seconds = float(item.get("seconds") or 0)
    elapsed = f"{seconds:.0f}s" if seconds < 90 else f"{seconds / 60:.0f}m"
    return f"{_KIND_ICONS.get(kind, '•')} {words}  {elapsed}"


# ----------------------------------------------------------------- live working view

def pack_above_expanding(widget: Any, **options: Any) -> None:
    """Pack a strip at the bottom of its parent so it actually gets height.

    Tk's packer gives space in packing order, and the chat area packed earlier
    with expand=True took all of it: a strip packed afterwards got one pixel and
    never appeared. Packing it just before the expanding sibling fixes that.
    """
    expanding = next((child for child in widget.master.pack_slaves()
                      if child is not widget and str(child.pack_info().get("expand")) in {"1", "True", "true"}), None)
    if expanding is not None:
        widget.pack(side="bottom", fill="x", before=expanding, **options)
    else:
        widget.pack(side="bottom", fill="x", **options)


class WorkingView(tk.Frame):
    """The steps of the running turn, shown as they happen, with what the screen looks like after each action.

    Nothing here is a model's private reasoning: each line is a real operation
    the runtime started (a tool, a model call, a search), ticked when it ends.
    """

    THUMB_WIDTH = 300

    def __init__(self, parent: Any, palette: Any) -> None:
        super().__init__(parent, bg=palette.PANEL_2, highlightthickness=1, highlightbackground=palette.LINE)
        self.palette = palette
        self.title = StringVar(self, value="Working")
        tk.Label(self, textvariable=self.title, bg=palette.PANEL_2, fg=palette.GOLD, anchor="w",
                 font=("Segoe UI Semibold", 10)).pack(fill="x", padx=8, pady=(5, 0))
        body = tk.Frame(self, bg=palette.PANEL_2)
        body.pack(fill="x", padx=8, pady=(2, 6))
        self.steps = tk.Text(body, height=5, wrap="word", state="disabled", bg=palette.PANEL_2,
                             fg=palette.WHITE, relief="flat", font=("Segoe UI", 9), highlightthickness=0)
        self.steps.pack(side="left", fill="both", expand=True)
        self.steps.tag_configure("done", foreground=palette.MUTED)
        self.steps.tag_configure("live", foreground=palette.WHITE)
        self.shot = tk.Label(body, bg=palette.PANEL_2)
        self._photo: Any = None
        self._seen: dict[str, dict[str, Any]] = {}
        self._order: list[str] = []
        self._started = 0.0
        self._active = False

    def destroy(self) -> None:
        self._photo = None
        super().destroy()

    def begin(self) -> None:
        self._seen.clear()
        self._order.clear()
        self._started = time.monotonic()
        self._active = True
        self.shot.pack_forget()
        self._photo = None
        self._paint()

    def end(self) -> None:
        self._active = False
        self.pack_forget()

    def update_from(self, running: list[dict[str, Any]]) -> list[str]:
        """Fold the current activity list in; returns the computer actions that just finished."""
        if not self._active:
            return []
        live_ids = set()
        for item in running:
            key = str(item.get("activity_id"))
            live_ids.add(key)
            if key not in self._seen:
                self._order.append(key)
            self._seen[key] = {**item, "finished": False}
        finished_actions = []
        for key, item in self._seen.items():
            if key not in live_ids and not item.get("finished"):
                item["finished"] = True
                if item.get("kind") == "tool" and str(item.get("label", "")).startswith(COMPUTER_ACTIONS):
                    finished_actions.append(str(item.get("label")))
        self._paint()
        return finished_actions

    def show_screenshot(self, image: Any, path: Any = None) -> None:
        """Show what the screen looked like after the last action; a click opens it full size."""
        if not self._active or image is None:
            return
        try:
            from PIL import ImageTk

            ratio = self.THUMB_WIDTH / max(1, image.width)
            thumb = image.resize((self.THUMB_WIDTH, max(1, int(image.height * ratio))))
            self._photo = ImageTk.PhotoImage(thumb, master=self)
        except Exception:
            return
        self.shot.configure(image=self._photo, cursor="hand2" if path else "")
        if path:
            from app.platform_support import open_path

            self.shot.bind("<Button-1>", lambda _event, target=path: open_path(target))
        if not self.shot.winfo_manager():
            # Before the expanding steps list, or it takes all the width first.
            self.shot.pack(side="right", padx=(8, 0), before=self.steps)

    def _paint(self) -> None:
        live = sum(1 for item in self._seen.values() if not item.get("finished"))
        took = time.monotonic() - self._started
        self.title.set(f"Working · {len(self._order)} step{'s' if len(self._order) != 1 else ''} · "
                       f"{live} running · {took:.0f}s")
        rows = []
        for key in self._order[-12:]:
            item = self._seen[key]
            mark = "✓ " if item.get("finished") else "… "
            rows.append((mark + activity_line(item) + "\n", "done" if item.get("finished") else "live"))
        if rows != getattr(self, "_painted_steps", None):
            view = self.steps.yview()
            follow = view[1] >= .98
            self.steps.configure(state="normal")
            self.steps.delete("1.0", "end")
            for line, tag in rows:
                self.steps.insert("end", line, tag)
            self.steps.configure(state="disabled")
            if follow:
                self.steps.see("end")
            else:
                self.steps.yview_moveto(view[0])
            self._painted_steps = rows
        if self._active and self._order and not self.winfo_manager():
            pack_above_expanding(self, pady=4)
