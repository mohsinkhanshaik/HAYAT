"""Reusable native starter kits and persistent product-workspace tools."""
from __future__ import annotations

import io
from uuid import uuid4

from app.tools.file_editing import FileEditing, atomic_save

PRESETS = {
    'weekly_planner': ('Weekly plan', ['Task', 'Project', 'Due date', 'Planned hours', 'Actual hours', 'Status'],
                       ['Plan priorities', 'Schedule realistic commitments', 'Review the week']),
    'content_calendar': ('Content calendar', ['Publish date', 'Channel', 'Audience question', 'Content angle', 'Call to action', 'Status'],
                         ['Choose an audience question', 'Draft and check the evidence', 'Review before publishing']),
    'research_report': ('Evidence log', ['Claim', 'Source URL', 'Retrieved date', 'Evidence', 'Confidence', 'Status'],
                        ['Define the research question', 'Collect attributable evidence', 'Separate findings from assumptions']),
}


class ProductWorkshop:
    def __init__(self, store, filesystem, business):
        self.store, self.filesystem, self.business = store, filesystem, business

    def build_starter(self, product_id: str, kind: str = 'weekly_planner') -> dict:
        if kind not in PRESETS:
            raise ValueError('Choose weekly_planner, content_calendar, or research_report.')
        product = self.store._product(product_id)
        profile = self.business.profile()
        folder = self.store.output.parent / 'product-workshop' / product_id / uuid4().hex
        folder = self.filesystem._path(str(folder), write=True)
        sheet_name, headings, steps = PRESETS[kind]
        brand = profile.get('business_name') or 'HAYAT'
        owner = profile.get('owner_name') or ''
        guide = f'# {product["title"]}\n\n{brand}' + (f' · Prepared by {owner}' if owner else '')
        guide += f'\n\n## Who this is for\n\n{product["audience"]}\n\n## Your brief\n\n{product["description"]}\n'
        for index, step in enumerate(steps, 1):
            guide += f'\n## {index}. {step}\n\n'
            guide += {
                'weekly_planner': [
                    'Write your three most useful outcomes for this week. Break each into tasks with a due date and estimated hours. Replace the example row before using the totals.',
                    'Compare planned work with available hours. Leave room for interruptions. Set one owner and one next action for each commitment.',
                    'Record actual hours and delivery status. Use the overview totals to compare the plan with completed work. Carry unfinished items into the next week deliberately.',
                ],
                'content_calendar': [
                    'Capture a specific question asked by your audience. Add the intended channel and a clear angle for each draft. Keep customer information out of public examples.',
                    'Draft the useful answer first. Check factual claims against attributable sources. Adapt the same idea to the selected channel and write one clear call to action.',
                    'Use Draft, Review, Ready, and Published to track the workflow. Mark Published only after you actually publish externally. Record observed results separately from forecasts.',
                ],
                'research_report': [
                    'Write a narrow question and explain the decision this report will support. Set the scope and a retrieval date before collecting sources.',
                    'Use one evidence-log row per claim. Keep the original source URL, retrieval date, supporting excerpt or observation, and confidence. Leave unsupported claims unresolved.',
                    'Write findings from the evidence log. Label assumptions, conflicting evidence, and missing information. The generated template contains no independently researched findings.',
                ],
            }[kind][index - 1] + '\n'
        guide += '\n## Files included\n\n- Editable workbook with an overview, working sheet, and instructions.\n- This guide in Word and PDF.\n- A review presentation to customize.\n'
        guide += '\n## Review before release\n\nThis is a starter template generated from your brief. Replace example content, verify claims, and review your product license and listing before selling it. Checkout is configured separately.\n'
        paths = [self.filesystem.write_file(str(folder / 'README.md'), guide)['path'],
                 self.filesystem.write_pdf(str(folder / 'guide.pdf'), guide)['path'],
                 self.filesystem.write_docx(str(folder / 'guide.docx'), guide)['path']]
        editor = FileEditing(self.filesystem)
        example = (['Example — replace before use', 'Your project', '', 0, 0, 'Planned'] if kind == 'weekly_planner'
                   else ['', 'Your channel', 'Example — replace with an actual question', '', '', 'Draft'] if kind == 'content_calendar'
                   else ['Example — replace with an evidenced claim', '', '', '', 'Unassessed', 'Review'])
        workbook = editor.write_spreadsheet(str(folder / 'workbook.xlsx'), [headings, example], sheet_name)
        self._style_workbook(workbook['path'], product, kind, brand, sheet_name)
        paths.append(workbook['path'])
        slides = [{'title': product['title'], 'body': f'{brand}\nFor {product["audience"]}\n{product["description"]}'}]
        slides.extend({'title': step, 'body': 'Evidence / completed work:\n\nWhat needs another step:\n\nOwner and next action:'} for step in steps)
        paths.append(editor.write_presentation(str(folder / 'review.pptx'), slides)['path'])
        self._style_slides(paths[-1])
        result = self.store.set_files(product_id, paths)
        return {**result, 'kind': kind, 'folder': str(folder), 'review_required': True,
                'status': 'Starter files created. Edit and review them, then build the delivery ZIP.'}

    def _style_workbook(self, path, product, kind, brand, sheet_name):
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.datavalidation import DataValidation

        book = openpyxl.load_workbook(path)
        sheet = book[sheet_name]
        overview = book.create_sheet('Overview', 0)
        overview.append([brand])
        overview.append([product['title']])
        overview.append(['Audience', product['audience']])
        overview.append(['Brief', product['description']])
        overview.append(['Status', 'Starter draft — replace the example row before using totals'])
        for row in overview.iter_rows(max_row=5):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'  # Supplied listing text is never a formula.
        if kind == 'weekly_planner':
            overview.append(['Planned hours', f"=SUM('{sheet_name}'!D2:D501)"])
            overview.append(['Actual hours', f"=SUM('{sheet_name}'!E2:E501)"])
            statuses = 'Planned,In progress,Done,Deferred'
        else:
            overview.append(['Ready items', f'=COUNTIF(\'{sheet_name}\'!F2:F501,"Ready")'])
            overview.append(['Items to review', f'=COUNTIF(\'{sheet_name}\'!F2:F501,"Review")'])
            statuses = 'Draft,Review,Ready,Published' if kind == 'content_calendar' else 'Review,Ready,Conflicting,Unresolved'
        instructions = book.create_sheet('Instructions')
        for row in [['Start here', 'How to use'], ['1', 'Replace the clearly labelled example row.'],
                    ['2', 'Add your work to rows 2–501. Overview formulas cover that range.'],
                    ['3', 'Choose statuses from the dropdown in column F.'],
                    ['4', 'Record evidence and actual results; forecasts are separate.']]:
            instructions.append(row)
        validation = DataValidation(type='list', formula1=f'"{statuses}"', allow_blank=True)
        sheet.add_data_validation(validation)
        validation.add('F2:F501')
        sheet.auto_filter.ref = 'A1:F501'
        for target in book:
            target.sheet_view.showGridLines = False
            target.freeze_panes = 'A2'
            for cell in target[1]:
                cell.fill = PatternFill('solid', fgColor='18263B')
                cell.font = Font(name='Calibri', size=12, bold=True, color='F0CF72')
            for row in target.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
                    cell.font = Font(name='Calibri', size=11, color='152033')
                    if cell.row % 2 == 0:
                        cell.fill = PatternFill('solid', fgColor='E8EDF5')
            for column in target.columns:
                target.column_dimensions[column[0].column_letter].width = 24 if target is sheet else 30
            target.row_dimensions[1].height = 28
        overview.column_dimensions['B'].width = 72
        overview.row_dimensions[4].height = 65
        instructions.column_dimensions['B'].width = 85
        output = io.BytesIO()
        book.save(output)
        book.close()
        target = self.filesystem._path(path, write=True)
        self.filesystem._ensure_capacity(len(output.getvalue()))
        atomic_save(target, output.getvalue())

    def _style_slides(self, path):
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.util import Pt

        deck = Presentation(path)
        for slide in deck.slides:
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor.from_string('101A2B')
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                for paragraph in shape.text_frame.paragraphs:
                    paragraph.font.name = 'Calibri'
                    paragraph.font.size = Pt(30 if shape == slide.shapes.title else 18)
                    paragraph.font.color.rgb = RGBColor.from_string('F0CF72' if shape == slide.shapes.title else 'F5F7FB')
        output = io.BytesIO()
        deck.save(output)
        target = self.filesystem._path(path, write=True)
        self.filesystem._ensure_capacity(len(output.getvalue()))
        atomic_save(target, output.getvalue())


def register_product_workspace_tools(registry, store, filesystem, business):
    workshop = ProductWorkshop(store, filesystem, business)
    text = {'type': 'string', 'minLength': 1, 'maxLength': 20000}
    product = {'product_id': text}
    fields = {key: text for key in ('product_id', 'title', 'audience', 'description', 'price', 'currency', 'license_text')}
    fields['expected_sha256'] = {'type': 'string', 'pattern': '^[a-f0-9]{64}$'}
    definitions = [
        ('product_details', 'Read a product draft, metadata SHA256, selected deliverables and latest tracked ZIP.', product, ['product_id'], store.details, False),
        ('product_update', 'Update a local product listing with a conflict check. Currency cannot change after recorded sales. Earlier ZIPs keep their original listing.', fields, [key for key in fields if key != 'currency'], store.update, True),
        ('product_set_files', 'Save the chosen product deliverable files for reuse across sessions. An empty list clears the selection. Does not package or publish.', {**product, 'files': {'type': 'array', 'minItems': 0, 'maxItems': 100, 'items': text}}, ['product_id', 'files'], store.set_files, True),
        ('product_verify_bundle', 'Recheck the latest tracked product ZIP, hashes and license, and detect source/listing changes requiring a rebuild. Does not verify factual or commercial quality.', product, ['product_id'], store.verify_bundle, False),
        ('product_build_starter', 'Create branded, editable native PDF/Word/Excel/PowerPoint starter files for an existing product. Choose weekly_planner, content_calendar or research_report. Examples require review; no research, publication or payment occurs.', {**product, 'kind': {'type': 'string', 'enum': list(PRESETS)}}, ['product_id'], workshop.build_starter, True),
    ]
    for name, description, properties, required, handler, write in definitions:
        registry.register(name, description, {'type': 'object', 'properties': properties, 'required': required, 'additionalProperties': False},
                          'medium' if write else 'low', 'filesystem_write_project' if write else 'filesystem_read', handler)
    return workshop
