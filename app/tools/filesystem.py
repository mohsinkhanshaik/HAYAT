from __future__ import annotations

import shutil
from pathlib import Path
from typing import TYPE_CHECKING, Any

from app.safety.permissions import PermissionManager
from app.tools.registry import ToolRegistry

if TYPE_CHECKING:
    from app.storage.manager import StorageManager

#: Resolved once: (body_font, bold_font, unicode_ok). The built-in Helvetica
#: only covers Latin-1, so a PDF of anything with a bullet or a smart quote came
#: out mangled. When a system TrueType font is present we register it and get
#: full Unicode; otherwise we keep Helvetica and sanitise to Latin-1 rather than
#: crash. None until the first PDF is written.
_PDF_FONTS: dict[str, Any] | None = None

#: Unicode TrueType families to try, in order, across operating systems. Each is
#: (regular, bold, italic, bold-italic); the first whose regular face exists is
#: registered, so a bare "make a pdf" renders real Unicode on Windows (Arial),
#: Linux (DejaVu or Liberation), and macOS (Arial under Supplemental) alike
#: rather than only on the machine it was first written on. A missing bold or
#: italic face falls back to the regular one.
_PDF_FONT_FAMILIES: tuple[tuple[str, str, str, str], ...] = (
    (r"C:\Windows\Fonts\arial.ttf", r"C:\Windows\Fonts\arialbd.ttf",
     r"C:\Windows\Fonts\ariali.ttf", r"C:\Windows\Fonts\arialbi.ttf"),
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
     "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
     "/usr/share/fonts/truetype/dejavu/DejaVuSans-Oblique.ttf",
     "/usr/share/fonts/truetype/dejavu/DejaVuSans-BoldOblique.ttf"),
    ("/usr/share/fonts/dejavu/DejaVuSans.ttf",
     "/usr/share/fonts/dejavu/DejaVuSans-Bold.ttf",
     "/usr/share/fonts/dejavu/DejaVuSans-Oblique.ttf",
     "/usr/share/fonts/dejavu/DejaVuSans-BoldOblique.ttf"),
    ("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
     "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
     "/usr/share/fonts/truetype/liberation/LiberationSans-Italic.ttf",
     "/usr/share/fonts/truetype/liberation/LiberationSans-BoldItalic.ttf"),
    ("/System/Library/Fonts/Supplemental/Arial.ttf",
     "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
     "/System/Library/Fonts/Supplemental/Arial Italic.ttf",
     "/System/Library/Fonts/Supplemental/Arial Bold Italic.ttf"),
    ("/Library/Fonts/Arial.ttf", "/Library/Fonts/Arial Bold.ttf",
     "/Library/Fonts/Arial Italic.ttf", "/Library/Fonts/Arial Bold Italic.ttf"),
)


def _select_font_family(exists: Any = None) -> tuple[str, str, str, str] | None:
    """The first candidate family whose regular face is present, or None.

    Pure and side-effect-free so the cross-platform choice can be tested with a
    fake existence check rather than the machine's actual fonts.
    """
    check = exists or (lambda path: Path(path).exists())
    for family in _PDF_FONT_FAMILIES:
        if check(family[0]):
            return family
    return None


def _pdf_fonts() -> tuple[str, str, bool]:
    """(body, bold, unicode_ok). Registers a full Unicode TrueType family --

    normal, bold, italic, bold-italic -- from the first available system font
    set, so **bold**/*italic* markup resolve to real weights and non-Latin-1
    text (bullets, symbols, accents) renders on any OS. Falls back to the
    built-in Helvetica family, whose four weights already exist, otherwise.
    """
    global _PDF_FONTS
    if _PDF_FONTS is None:
        _PDF_FONTS = {}
        try:
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont

            family = _select_font_family()
            if family is not None:
                regular, bold, italic, bold_italic = family
                faces = {"HAYATBody": regular, "HAYATBold": bold,
                         "HAYATItalic": italic, "HAYATBoldItalic": bold_italic}
                for name, candidate in faces.items():
                    pdfmetrics.registerFont(TTFont(
                        name, candidate if Path(candidate).exists() else regular))
                pdfmetrics.registerFontFamily(
                    "HAYATBody", normal="HAYATBody", bold="HAYATBold",
                    italic="HAYATItalic", boldItalic="HAYATBoldItalic")
                _PDF_FONTS = {"body": "HAYATBody", "bold": "HAYATBold", "unicode": True}
        except Exception:
            _PDF_FONTS = {}
    if _PDF_FONTS.get("unicode"):
        return _PDF_FONTS["body"], _PDF_FONTS["bold"], True
    return "Helvetica", "Helvetica-Bold", False


def _md_inline(text: str) -> str:
    """Turn a line of light markdown into reportlab's inline markup.

    Escapes the XML-special characters first so a literal < in the user's text
    is not read as a tag, then maps **bold**, *italic* and `code` onto the
    <b>/<i>/<font> that Paragraph understands.
    """
    import html
    import re

    text = html.escape(text, quote=False)
    text = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)
    text = re.sub(r"(?<![*\w])\*(?!\s)(.+?)(?<!\s)\*(?!\*)", r"<i>\1</i>", text)
    text = re.sub(r"`(.+?)`", r'<font face="Courier">\1</font>', text)
    return text


def _md_to_flowables(content: str, title: str | None, styles: dict[str, Any]) -> list[Any]:
    """Lay light markdown out as a real document, not a wall of lines.

    Headings become headings, ``-``/``*`` runs become a bulleted list, ``1.``
    runs become a numbered list, blank lines separate paragraphs, and
    consecutive plain lines join into one flowing paragraph that wraps to the
    page. This is what makes a generated PDF read like a document a person can
    follow rather than a screen dump.
    """
    import re

    from reportlab.platypus import ListFlowable, ListItem, Paragraph, Spacer

    flow: list[Any] = []
    if title:
        flow.append(Paragraph(_md_inline(title), styles["Title"]))
        flow.append(Spacer(1, 10))

    paragraph: list[str] = []
    bullets: list[str] = []
    numbers: list[str] = []

    def flush_paragraph() -> None:
        if paragraph:
            flow.append(Paragraph(" ".join(paragraph), styles["Body"]))
            paragraph.clear()

    def flush_bullets() -> None:
        if bullets:
            flow.append(ListFlowable(
                [ListItem(Paragraph(item, styles["Body"]), leftIndent=14) for item in bullets],
                bulletType="bullet", start="\u2022", leftIndent=10))
            bullets.clear()

    def flush_numbers() -> None:
        if numbers:
            flow.append(ListFlowable(
                [ListItem(Paragraph(item, styles["Body"]), leftIndent=14) for item in numbers],
                bulletType="1", leftIndent=10))
            numbers.clear()

    def flush_all() -> None:
        flush_paragraph()
        flush_bullets()
        flush_numbers()

    for line in content.split("\n"):
        stripped = line.strip()
        if not stripped:
            flush_all()
            continue
        heading = re.match(r"(#{1,4})\s+(.*)", stripped)
        bullet = re.match(r"[-*+]\s+(.*)", stripped)
        number = re.match(r"\d+[.)]\s+(.*)", stripped)
        if heading:
            flush_all()
            level = min(len(heading.group(1)), 3)
            flow.append(Paragraph(_md_inline(heading.group(2)), styles[f"H{level}"]))
        elif bullet:
            flush_paragraph()
            flush_numbers()
            bullets.append(_md_inline(bullet.group(1)))
        elif number:
            flush_paragraph()
            flush_bullets()
            numbers.append(_md_inline(number.group(1)))
        else:
            flush_bullets()
            flush_numbers()
            paragraph.append(_md_inline(stripped))
    flush_all()

    if not flow:
        from reportlab.platypus import Paragraph as _P
        flow.append(_P("&nbsp;", styles["Body"]))
    return flow


def _docx_add_runs(paragraph: Any, text: str) -> None:
    """Add text to a docx paragraph, honouring **bold**, *italic*, `code`."""
    import re

    pattern = re.compile(r"\*\*(.+?)\*\*|(?<!\*)\*(?!\s)(.+?)(?<!\s)\*(?!\*)|`(.+?)`")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        bold, italic, code = match.group(1), match.group(2), match.group(3)
        if bold is not None:
            paragraph.add_run(bold).bold = True
        elif italic is not None:
            paragraph.add_run(italic).italic = True
        elif code is not None:
            run = paragraph.add_run(code)
            run.font.name = "Consolas"
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _md_to_docx(document: Any, content: str) -> None:
    """Lay light markdown out as a Word document: headings, bulleted and

    numbered lists, wrapped paragraphs, and inline bold/italic -- the same
    structure the PDF path produces, so the two formats read alike.
    """
    import re

    paragraph: list[str] = []

    def flush() -> None:
        if paragraph:
            _docx_add_runs(document.add_paragraph(), " ".join(paragraph))
            paragraph.clear()

    for line in content.split("\n"):
        stripped = line.strip()
        if not stripped:
            flush()
            continue
        heading = re.match(r"(#{1,4})\s+(.*)", stripped)
        bullet = re.match(r"[-*+]\s+(.*)", stripped)
        number = re.match(r"\d+[.)]\s+(.*)", stripped)
        if heading:
            flush()
            plain = re.sub(r"[*_`]", "", heading.group(2))
            document.add_heading(plain, level=min(len(heading.group(1)), 4))
        elif bullet:
            flush()
            _docx_add_runs(document.add_paragraph(style="List Bullet"), bullet.group(1))
        elif number:
            flush()
            _docx_add_runs(document.add_paragraph(style="List Number"), number.group(1))
        else:
            paragraph.append(stripped)
    flush()


class FilesystemTools:
    IGNORED_DIRECTORY_NAMES = {".git", ".venv", "__pycache__", ".pytest_cache", ".mypy_cache", ".ruff_cache", "node_modules"}
    MAX_LIST_ENTRIES = 2_000
    MAX_SEARCH_MATCHES = 500
    MAX_READ_CHARS = 1_000_000

    def __init__(self, permissions: PermissionManager, storage: "StorageManager | None" = None) -> None:
        self.permissions = permissions
        self.storage = storage

    def _path(self, path: str, write: bool = False) -> Path:
        return self.permissions.validate_project_path(path, write=write)

    #: Bounds for finding a file by name: never a slow crawl of the whole disk.
    FIND_SECONDS, FIND_FILES = 2.0, 60_000

    def find(self, name: str, want_folder: bool = False) -> list[Path]:
        """Places a bare or misplaced name ("settings.yaml", "notes/plan.md") exists inside the allowed roots."""
        import os
        import time
        wanted = Path(str(name).replace("\\", "/"))
        parts = [part.lower() for part in wanted.parts if part not in (".", "")]
        if not parts:
            return []
        started, seen, matches = time.monotonic(), 0, []
        for root in self.permissions.allowed_roots():
            if not root.is_dir():
                continue
            for folder, directories, files in os.walk(root):
                directories[:] = [item for item in directories if item not in self.IGNORED_DIRECTORY_NAMES and not item.startswith(".")]
                seen += len(files)
                for item in directories if want_folder else files:
                    if item.lower() != parts[-1]:
                        continue
                    candidate = Path(folder) / item
                    tail = [part.lower() for part in candidate.parts[-len(parts):]]
                    if tail == parts:
                        matches.append(candidate)
                if seen > self.FIND_FILES or time.monotonic() - started > self.FIND_SECONDS or len(matches) > 10:
                    break
        unique = []
        for match in matches:
            if match not in unique:
                unique.append(match)
        return unique[:10]

    def _existing(self, path: str, want_folder: bool = False) -> tuple[Path, str]:
        """The path, or the one place a file of that name exists; several places are offered, never guessed."""
        try:
            resolved = self._path(path)
        except FileNotFoundError:
            resolved = None
        if resolved is not None and resolved.exists():
            return resolved, ""
        found = self.find(path, want_folder)
        if len(found) == 1:
            located = self._path(str(found[0]))
            return located, f"{path} was not where it was asked for; found it at {self._display_path(located)}"
        kind = "folder" if want_folder else "file"
        if found:
            raise FileNotFoundError(f"No {kind} at {path}. Did you mean: " + "; ".join(self._display_path(item) for item in found[:6]) + "?")
        raise FileNotFoundError(f"No {kind} named {Path(path).name} in {path} or anywhere in the allowed folders.")

    def _containing_root(self, resolved: Path) -> Path:
        roots = self.permissions.allowed_roots()
        return next(
            (root for root in roots if resolved == root or root in resolved.parents),
            self.permissions.project_root,
        )

    def _display_path(self, item: Path) -> str:
        try:
            return item.relative_to(self._containing_root(item)).as_posix()
        except ValueError:
            return str(item)

    def _ignored(self, item: Path) -> bool:
        return any(part in self.IGNORED_DIRECTORY_NAMES for part in Path(self._display_path(item)).parts)

    def list_folder(self, path: str = ".", recursive: bool = False, max_entries: int = MAX_LIST_ENTRIES) -> dict[str, Any]:
        folder, note = self._existing(path, want_folder=True)
        if not folder.is_dir():
            raise NotADirectoryError(f"{self._display_path(folder)} is a file, not a folder: read it with read_file instead.")
        limit = min(max(1, int(max_entries)), 10_000)
        iterator = folder.rglob("*") if recursive else folder.iterdir()
        entries = []
        truncated = False
        for item in iterator:
            if self._ignored(item):
                continue
            entries.append({
                "path": self._display_path(item),
                "type": "directory" if item.is_dir() else "file",
                "size": item.stat().st_size if item.is_file() else None,
            })
            if len(entries) >= limit:
                truncated = True
                break
        entries.sort(key=lambda value: value["path"].lower())
        return {"path": str(folder), "entries": entries, "count": len(entries), "truncated": truncated, **({"note": note} if note else {})}

    def read_file(self, path: str, max_chars: int = MAX_READ_CHARS) -> dict[str, Any]:
        file_path, note = self._existing(path)
        if file_path.is_dir():
            raise IsADirectoryError(f"{self._display_path(file_path)} is a folder: list it with list_folder instead.")
        limit = min(max(1, int(max_chars)), self.MAX_READ_CHARS)
        chunks: list[str] = []
        total = 0
        truncated = False
        with file_path.open("r", encoding="utf-8", errors="replace") as handle:
            while total < limit:
                piece = handle.read(min(65_536, limit - total))
                if not piece:
                    break
                chunks.append(piece)
                total += len(piece)
            truncated = bool(handle.read(1))
        content = "".join(chunks)
        return {"path": str(file_path), "content": content, "truncated": truncated, "size": file_path.stat().st_size,
                **({"note": note} if note else {})}

    def write_file(self, path: str, content: str) -> dict[str, Any]:
        # A document format must be a real document, not text with a document's
        # extension. HAYAT asked to "make a pdf" once wrote the raw markdown into
        # a file named *.pdf; Adobe called it damaged because it was never a PDF.
        # Routing here guarantees a .pdf is a PDF and a .docx is a Word file no
        # matter which write tool the planner reached for.
        suffix = Path(path).suffix.lower()
        if suffix == ".pdf":
            return self.write_pdf(path, content)
        if suffix == ".docx":
            return self.write_docx(path, content)
        if suffix in {".xlsx", ".pptx", ".zip", ".png", ".jpg", ".mp4", ".db", ".exe", ".dll"}:
            raise ValueError("Use a format-specific tool to create or edit this binary format.")
        file_path = self._path(path, write=True)
        encoded = content.encode("utf-8")
        previous = file_path.stat().st_size if file_path.is_file() else 0
        self._ensure_capacity(max(0, len(encoded) - previous))
        file_path.parent.mkdir(parents=True, exist_ok=True)
        from app.tools.file_editing import atomic_save
        self._ensure_capacity(len(encoded))
        return atomic_save(file_path, encoded)

    def append_file(self, path: str, content: str) -> dict[str, Any]:
        file_path = self._path(path, write=True)
        from app.tools.file_editing import FileEditing
        editor = FileEditing(self)
        if not file_path.exists():
            return self.write_file(path, content)
        current = editor.read_editable_file(path)
        return editor.save_edited_file(path, current["content"] + content, current["sha256"])

    def create_folder(self, path: str) -> dict[str, Any]:
        folder = self._path(path, write=True)
        folder.mkdir(parents=True, exist_ok=True)
        return {"path": str(folder), "exists": folder.is_dir()}

    def _stored_path(self, path: str, suffix: str) -> str:
        """Force the suffix, and send a bare filename to the artifacts folder.

        "make a pdf called plan" should land in data/artifacts/plan.pdf, not the
        project root, so generated documents do not pile up beside the source.
        A path that already names a directory is left where the caller put it.
        """
        candidate = Path(path)
        if candidate.suffix.lower() != suffix:
            candidate = candidate.with_suffix(suffix)
        if candidate.parent in (Path("."), Path("")):
            candidate = Path("data") / "artifacts" / candidate.name
        return candidate.as_posix()

    def write_pdf(self, path: str, content: str, title: str | None = None) -> dict[str, Any]:
        """Render text/markdown into a formatted, readable PDF.

        HAYAT could read PDFs but never write one, and the first pass only
        stamped raw lines onto the page. This lays the content out as a real
        document: headings, bulleted and numbered lists, bold/italic, wrapped
        paragraphs, and a page number in the footer. The result is meant to be
        opened and read, not just to exist. Returns where it landed and how many
        pages it took so the reply can name a file the user can actually open.
        """
        from reportlab.lib.pagesizes import LETTER
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import inch
        from reportlab.platypus import SimpleDocTemplate
        import io
        from app.tools.file_editing import atomic_save

        file_path = self._path(self._stored_path(path, ".pdf"), write=True)
        self._ensure_capacity(len(content.encode("utf-8")) + 8192)
        file_path.parent.mkdir(parents=True, exist_ok=True)

        body_font, bold_font, unicode_ok = _pdf_fonts()
        clean = (str) if unicode_ok else (
            lambda value: value.encode("latin-1", "replace").decode("latin-1"))

        styles = {
            "Title": ParagraphStyle("Title", fontName=bold_font, fontSize=20, leading=25, spaceAfter=6),
            "H1": ParagraphStyle("H1", fontName=bold_font, fontSize=16, leading=20, spaceBefore=12, spaceAfter=6),
            "H2": ParagraphStyle("H2", fontName=bold_font, fontSize=13, leading=17, spaceBefore=10, spaceAfter=5),
            "H3": ParagraphStyle("H3", fontName=bold_font, fontSize=11.5, leading=15, spaceBefore=8, spaceAfter=4),
            "Body": ParagraphStyle("Body", fontName=body_font, fontSize=11, leading=15, spaceAfter=6),
        }
        flow = _md_to_flowables(clean(content), clean(title) if title else None, styles)

        def footer(canvas: Any, document: Any) -> None:
            canvas.setFont(body_font, 9)
            canvas.drawCentredString(LETTER[0] / 2, 0.55 * inch, str(document.page))

        output = io.BytesIO()
        document = SimpleDocTemplate(
            output, pagesize=LETTER, title=clean(title) if title else file_path.stem,
            topMargin=inch, bottomMargin=inch, leftMargin=inch, rightMargin=inch)
        document.build(flow, onFirstPage=footer, onLaterPages=footer)
        self._ensure_capacity(len(output.getvalue()))
        return {**atomic_save(file_path, output.getvalue()), "pages": document.page, "format": "pdf"}

    def write_docx(self, path: str, content: str, title: str | None = None) -> dict[str, Any]:
        """Render text/markdown into an editable Word (.docx) document.

        The companion to write_pdf for when the file needs to be changed rather
        than just read: same markdown structure -- headings, lists, bold/italic,
        wrapped paragraphs -- laid into a real .docx that Word, LibreOffice, and
        Google Docs open natively. A bare name goes to the artifacts folder.
        """
        from docx import Document
        import io
        from app.tools.file_editing import atomic_save

        file_path = self._path(self._stored_path(path, ".docx"), write=True)
        self._ensure_capacity(len(content.encode("utf-8")) + 8192)
        file_path.parent.mkdir(parents=True, exist_ok=True)

        document = Document()
        if title:
            document.add_heading(title, level=0)
        _md_to_docx(document, content)
        output = io.BytesIO()
        document.save(output)
        self._ensure_capacity(len(output.getvalue()))
        return {**atomic_save(file_path, output.getvalue()), "format": "docx"}

    def copy_file(self, source: str, destination: str) -> dict[str, Any]:
        source_path = self._path(source)
        destination_path = self._path(destination, write=True)
        previous = destination_path.stat().st_size if destination_path.is_file() else 0
        self._ensure_capacity(max(0, source_path.stat().st_size - previous))
        destination_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source_path, destination_path)
        return {"source": str(source_path), "destination": str(destination_path), "exists": destination_path.exists()}

    def move_file(self, source: str, destination: str) -> dict[str, Any]:
        source_path = self._path(source, write=True)
        destination_path = self._path(destination, write=True)
        destination_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(source_path), str(destination_path))
        return {"source": str(source_path), "destination": str(destination_path), "exists": destination_path.exists()}

    def search_files(self, path: str = ".", pattern: str = "*", contains: str | None = None, max_matches: int = MAX_SEARCH_MATCHES) -> dict[str, Any]:
        folder = self._path(path)
        limit = min(max(1, int(max_matches)), 10_000)
        matches = []
        truncated = False
        for item in folder.rglob(pattern):
            if not item.is_file() or self._ignored(item):
                continue
            if contains is not None:
                try:
                    with item.open("r", encoding="utf-8", errors="ignore") as handle:
                        text = handle.read(self.MAX_READ_CHARS)
                    if contains.lower() not in text.lower():
                        continue
                except OSError:
                    continue
            matches.append(self._display_path(item))
            if len(matches) >= limit:
                truncated = True
                break
        matches.sort()
        return {"path": str(folder), "pattern": pattern, "matches": matches, "count": len(matches), "truncated": truncated}

    def file_exists(self, path: str) -> dict[str, Any]:
        candidate = self._path(path)
        return {"path": str(candidate), "exists": candidate.exists(), "is_file": candidate.is_file(), "is_directory": candidate.is_dir()}

    def get_file_info(self, path: str) -> dict[str, Any]:
        candidate = self._path(path)
        stat = candidate.stat()
        result = {"path": str(candidate), "size": stat.st_size, "modified": stat.st_mtime, "is_file": candidate.is_file(), "is_directory": candidate.is_dir()}
        if candidate.is_file() and stat.st_size <= 50 * 1024 * 1024:
            import hashlib
            result["sha256"] = hashlib.sha256(candidate.read_bytes()).hexdigest()
        return result

    def delete_file(self, path: str) -> dict[str, Any]:
        candidate = self._path(path, write=True)
        if not candidate.is_file():
            raise FileNotFoundError(candidate)
        candidate.unlink()
        return {"path": str(candidate), "deleted": not candidate.exists()}

    def _ensure_capacity(self, additional_bytes: int) -> None:
        if self.storage is not None:
            self.storage.ensure_capacity(additional_bytes)


def register_filesystem_tools(registry: ToolRegistry, storage: "StorageManager | None" = None) -> FilesystemTools:
    tools = FilesystemTools(registry.permissions, storage)
    object_schema = lambda properties, required=None: {"type": "object", "properties": properties, "required": required or [], "additionalProperties": False}
    registry.register("list_folder", "List files and folders inside an allowed root.", object_schema({"path": {"type": "string"}, "recursive": {"type": "boolean"}, "max_entries": {"type": "integer", "minimum": 1, "maximum": 10000}}), "low", "filesystem_read", tools.list_folder)
    registry.register("read_file", "Read a UTF-8 text file inside an allowed root.", object_schema({"path": {"type": "string"}, "max_chars": {"type": "integer", "minimum": 1, "maximum": 1000000}}, ["path"]), "low", "filesystem_read", tools.read_file)
    registry.register("write_file", "Write a UTF-8 file inside an allowed root, creating parent folders.", object_schema({"path": {"type": "string"}, "content": {"type": "string"}}, ["path", "content"]), "medium", "filesystem_write_project", tools.write_file)
    registry.register("append_file", "Append UTF-8 text to a file in an allowed root.", object_schema({"path": {"type": "string"}, "content": {"type": "string"}}, ["path", "content"]), "medium", "filesystem_write_project", tools.append_file)
    registry.register("create_folder", "Create a folder inside an allowed root.", object_schema({"path": {"type": "string"}}, ["path"]), "low", "filesystem_write_project", tools.create_folder)
    registry.register("write_pdf", "Render text content into a PDF document inside an allowed root, creating parent folders. Use for any 'make/save/export a PDF' request.", object_schema({"path": {"type": "string"}, "content": {"type": "string"}, "title": {"type": ["string", "null"]}}, ["path", "content"]), "medium", "filesystem_write_project", tools.write_pdf)
    registry.register("write_docx", "Render text content into an editable Word (.docx) document inside an allowed root, creating parent folders. Use for any 'make/save/export a Word document' request.", object_schema({"path": {"type": "string"}, "content": {"type": "string"}, "title": {"type": ["string", "null"]}}, ["path", "content"]), "medium", "filesystem_write_project", tools.write_docx)
    registry.register("copy_file", "Copy a file within allowed roots.", object_schema({"source": {"type": "string"}, "destination": {"type": "string"}}, ["source", "destination"]), "medium", "filesystem_write_project", tools.copy_file)
    registry.register("move_file", "Move a file within allowed roots.", object_schema({"source": {"type": "string"}, "destination": {"type": "string"}}, ["source", "destination"]), "medium", "filesystem_write_project", tools.move_file)
    registry.register("search_files", "Find files by glob and optional text inside an allowed root.", object_schema({"path": {"type": "string"}, "pattern": {"type": "string"}, "contains": {"type": ["string", "null"]}, "max_matches": {"type": "integer", "minimum": 1, "maximum": 10000}}), "low", "filesystem_read", tools.search_files)
    registry.register("file_exists", "Check whether a path exists inside an allowed root.", object_schema({"path": {"type": "string"}}, ["path"]), "low", "filesystem_read", tools.file_exists)
    registry.register("get_file_info", "Return size and timestamps for a path inside an allowed root.", object_schema({"path": {"type": "string"}}, ["path"]), "low", "filesystem_read", tools.get_file_info)
    registry.register("delete_file", "Delete one file inside an allowed root. Authority still applies.", object_schema({"path": {"type": "string"}}, ["path"]), "high", "filesystem_delete", tools.delete_file)
    from app.tools.file_editing import register_editing_tools
    register_editing_tools(registry, tools)
    return tools
