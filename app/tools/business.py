"""Freelance business toolkit: clients, time, quotes, invoices, payments,
expenses, work leads, an income goal, and a profit report.

Everything is local bookkeeping. HAYAT never collects money, logs into a
payment account, or sends anything from here: invoices and quotes are PDF
files the owner sends (email_send keeps its own approval), payments and
expenses are what the owner reports, and the report says so. Amounts are
integer cents in two-decimal currencies, the same rule the product ledger
uses, so the two can be summed in one profit report.
"""
from __future__ import annotations

import re
import json
import sqlite3
from datetime import date, datetime, timedelta, UTC
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from app.tools.products import CURRENCIES, cents
from app.tools.registry import ToolRegistry


LEAD_STATUSES = ("new", "contacted", "proposal", "won", "lost")
BASIS = ("Owner-reported records. Income is cash received in the period; profit is before tax. "
         "HAYAT has not verified any payment.")


def money(value: int, currency: str) -> str:
    sign = "-" if value < 0 else ""
    return f"{sign}{abs(value) / 100:,.2f} {currency}"


def _currency(value: str | None, default: str) -> str:
    currency = (value or default or "USD").strip().upper()
    if currency not in CURRENCIES:
        raise ValueError("Supported two-decimal currencies: " + ", ".join(sorted(CURRENCIES)) + ".")
    return currency


def _day(value: str | None) -> str:
    if not value or not str(value).strip():
        return date.today().isoformat()
    text = str(value).strip().lower()
    if text == "today":
        return date.today().isoformat()
    if text == "yesterday":
        return (date.today() - timedelta(days=1)).isoformat()
    try:
        return date.fromisoformat(text).isoformat()
    except ValueError as exc:
        raise ValueError("Dates must be YYYY-MM-DD, 'today' or 'yesterday'.") from exc


def _decimal(value: Any, what: str, low: str = "0", high: str = "100000") -> Decimal:
    try:
        number = Decimal(str(value).strip())
    except InvalidOperation as exc:
        raise ValueError(f"{what} must be a number.") from exc
    if not number.is_finite() or number < Decimal(low) or number > Decimal(high):
        raise ValueError(f"{what} must be between {low} and {high}.")
    return number


def period_bounds(period: str | None) -> tuple[str, str, str]:
    """(start, end exclusive, label) for 'this month', 'last month', 'this year', 'YYYY-MM', 'YYYY', 'all'."""
    text = (period or "this month").strip().lower()
    today = date.today()
    if text in {"", "this month", "month", "current month"}:
        text = today.strftime("%Y-%m")
    elif text in {"last month", "previous month"}:
        text = (today.replace(day=1) - timedelta(days=1)).strftime("%Y-%m")
    elif text in {"this year", "year", "current year"}:
        text = str(today.year)
    elif text in {"last year", "previous year"}:
        text = str(today.year - 1)
    if text in {"all", "all time", "everything"}:
        return "0000-01-01", "9999-12-31", "all time"
    if re.fullmatch(r"\d{4}-\d{2}", text):
        start = date.fromisoformat(text + "-01")
        end = (start.replace(day=28) + timedelta(days=4)).replace(day=1)
        return start.isoformat(), end.isoformat(), start.strftime("%B %Y")
    if re.fullmatch(r"\d{4}", text):
        return f"{text}-01-01", f"{int(text) + 1}-01-01", text
    raise ValueError("Period must be 'this month', 'last month', 'this year', 'last year', YYYY-MM, YYYY or 'all'.")


class BusinessLedger:
    """SQLite-backed books for a one-person business."""

    def __init__(self, path: Path, output: Path, permissions: Any, products: Any = None):
        self.path, self.output, self.permissions, self.products = Path(path), Path(output), permissions, products

    # ------------------------------------------------------------------ storage
    def _db(self) -> sqlite3.Connection:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        db = sqlite3.connect(self.path, timeout=15)
        db.row_factory = sqlite3.Row
        db.execute("PRAGMA foreign_keys=ON")
        db.executescript("""
            CREATE TABLE IF NOT EXISTS profile (key TEXT PRIMARY KEY, value TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS clients (
                id INTEGER PRIMARY KEY, name TEXT NOT NULL UNIQUE COLLATE NOCASE, email TEXT NOT NULL DEFAULT '',
                rate_cents INTEGER NOT NULL DEFAULT 0, currency TEXT NOT NULL, notes TEXT NOT NULL DEFAULT '',
                created TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS time_entries (
                id INTEGER PRIMARY KEY, client_id INTEGER NOT NULL REFERENCES clients(id),
                minutes INTEGER NOT NULL, description TEXT NOT NULL, work_date TEXT NOT NULL,
                rate_cents INTEGER NOT NULL, invoice_number TEXT, running_since TEXT, created TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS documents (
                number TEXT PRIMARY KEY, kind TEXT NOT NULL, client_id INTEGER NOT NULL REFERENCES clients(id),
                currency TEXT NOT NULL, issued TEXT NOT NULL, due TEXT NOT NULL,
                subtotal_cents INTEGER NOT NULL, tax_bp INTEGER NOT NULL, tax_cents INTEGER NOT NULL,
                total_cents INTEGER NOT NULL, status TEXT NOT NULL, notes TEXT NOT NULL DEFAULT '',
                path TEXT NOT NULL DEFAULT '', source_quote TEXT, created TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS document_items (
                number TEXT NOT NULL REFERENCES documents(number), line INTEGER NOT NULL,
                description TEXT NOT NULL, quantity TEXT NOT NULL, unit_cents INTEGER NOT NULL,
                amount_cents INTEGER NOT NULL, PRIMARY KEY (number, line));
            CREATE TABLE IF NOT EXISTS payments (
                id INTEGER PRIMARY KEY, number TEXT NOT NULL REFERENCES documents(number),
                amount_cents INTEGER NOT NULL, paid_on TEXT NOT NULL, reference TEXT UNIQUE, recorded TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS expenses (
                id INTEGER PRIMARY KEY, amount_cents INTEGER NOT NULL, currency TEXT NOT NULL,
                category TEXT NOT NULL, description TEXT NOT NULL, spent_on TEXT NOT NULL,
                reference TEXT UNIQUE, recorded TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS leads (
                id INTEGER PRIMARY KEY, title TEXT NOT NULL, source TEXT NOT NULL, url TEXT NOT NULL DEFAULT '',
                value_cents INTEGER NOT NULL DEFAULT 0, currency TEXT NOT NULL, status TEXT NOT NULL,
                follow_up TEXT, notes TEXT NOT NULL DEFAULT '', created TEXT NOT NULL, updated TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS other_income (
                id INTEGER PRIMARY KEY, amount_cents INTEGER NOT NULL, currency TEXT NOT NULL, source TEXT NOT NULL,
                description TEXT NOT NULL, received_on TEXT NOT NULL, reference TEXT UNIQUE, recorded TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS withdrawals (
                id INTEGER PRIMARY KEY, amount_cents INTEGER NOT NULL, currency TEXT NOT NULL,
                destination TEXT NOT NULL, status TEXT NOT NULL, note TEXT NOT NULL DEFAULT '', reference TEXT,
                requested_on TEXT NOT NULL, completed_on TEXT, recorded TEXT NOT NULL);
            -- The owner's own accounts, cards and payment handles: a nickname and the last 4 digits, never a full number.
            CREATE TABLE IF NOT EXISTS money_methods (
                id INTEGER PRIMARY KEY, kind TEXT NOT NULL, nickname TEXT NOT NULL, institution TEXT NOT NULL DEFAULT '',
                last4 TEXT NOT NULL DEFAULT '', currency TEXT NOT NULL, receive TEXT NOT NULL DEFAULT '',
                created TEXT NOT NULL, removed TEXT);
            -- Money the owner put into the business themselves: recorded, not moved, and never counted as income.
            CREATE TABLE IF NOT EXISTS funding (
                id INTEGER PRIMARY KEY, amount_cents INTEGER NOT NULL, currency TEXT NOT NULL, method_id INTEGER,
                note TEXT NOT NULL DEFAULT '', added_on TEXT NOT NULL, reference TEXT UNIQUE, recorded TEXT NOT NULL);
        """)
        return db

    def _run(self, work):
        db = self._db()
        try:
            with db:
                return work(db)
        finally:
            db.close()

    @staticmethod
    def _now() -> str:
        return datetime.now(UTC).isoformat(timespec="seconds")

    # ------------------------------------------------------------------ profile
    PROFILE_KEYS = ("business_name", "owner_name", "email", "address", "payment_instructions", "default_currency")

    def profile(self) -> dict:
        rows = self._run(lambda db: db.execute("SELECT key, value FROM profile").fetchall())
        values = {key: "" for key in self.PROFILE_KEYS}
        values.update({row["key"]: row["value"] for row in rows if row["key"] in self.PROFILE_KEYS})
        values["default_currency"] = values["default_currency"] or "USD"
        return values

    def set_profile(self, **fields: str) -> dict:
        updates = {key: str(value).strip() for key, value in fields.items() if value is not None and key in self.PROFILE_KEYS}
        if "default_currency" in updates:
            updates["default_currency"] = _currency(updates["default_currency"], "USD")
        if not updates:
            raise ValueError("Give at least one of: " + ", ".join(self.PROFILE_KEYS) + ".")
        self._run(lambda db: db.executemany("INSERT INTO profile VALUES (?, ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
                                            list(updates.items())))
        return {"saved": sorted(updates), "profile": self.profile()}

    def _default_currency(self) -> str:
        return self.profile()["default_currency"]

    def business_model(self, assumptions: dict | None = None) -> dict:
        """Read the saved plan, or evaluate an unsaved what-if without changing books."""
        from app.business_model import evaluate
        if assumptions is not None and not isinstance(assumptions, dict):
            raise ValueError("Business-model assumptions must be an object.")
        row = self._run(lambda db: db.execute("SELECT value FROM profile WHERE key='business_model'").fetchone())
        saved = json.loads(row["value"]) if row else {"currency": self._default_currency()}
        return evaluate({**saved, **(assumptions or {})})

    def save_business_model(self, assumptions: dict) -> dict:
        result = self.business_model(assumptions)
        self._run(lambda db: db.execute(
            "INSERT INTO profile VALUES ('business_model', ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
            (json.dumps(result["assumptions"], ensure_ascii=False),)))
        return {"saved": True, **result}

    def business_pilot(self, observations: dict | None = None, assumptions: dict | None = None) -> dict:
        """Read the last frozen pilot check or compare an unsaved cohort with the plan."""
        from app.business_pilot import evaluate_pilot
        if observations is not None:
            return evaluate_pilot(self.business_model(assumptions)['assumptions'], observations)
        if assumptions is not None:
            raise ValueError('Supply observations when comparing a different plan.')
        row = self._run(lambda db: db.execute("SELECT value FROM profile WHERE key='business_pilot'").fetchone())
        return json.loads(row['value']) if row else {'available': False}

    def save_business_pilot(self, observations: dict, assumptions: dict | None = None) -> dict:
        if not isinstance(observations, dict):
            raise ValueError('Pilot observations must be an object.')
        result = self.business_pilot(observations, assumptions)
        result.update(available=True, recorded_at=self._now())
        self._run(lambda db: db.execute(
            "INSERT INTO profile VALUES ('business_pilot', ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
            (json.dumps(result, ensure_ascii=False),)))
        return {'saved': True, **result}

    # ------------------------------------------------------------------ clients
    def add_client(self, name: str, email: str = "", hourly_rate: str = "0", currency: str | None = None,
                   notes: str = "") -> dict:
        name = (name or "").strip()
        if not name:
            raise ValueError("A client name is required.")
        record = {"name": name, "email": (email or "").strip(), "rate_cents": cents(hourly_rate or "0"),
                  "currency": _currency(currency, self._default_currency()), "notes": (notes or "").strip(),
                  "created": self._now()}

        def work(db):
            existing = db.execute("SELECT * FROM clients WHERE name=?", (name,)).fetchone()
            if existing:
                db.execute("UPDATE clients SET email=COALESCE(NULLIF(?, ''), email), rate_cents=CASE WHEN ?>0 THEN ? ELSE rate_cents END, "
                           "currency=?, notes=COALESCE(NULLIF(?, ''), notes) WHERE id=?",
                           (record["email"], record["rate_cents"], record["rate_cents"], record["currency"],
                            record["notes"], existing["id"]))
                return {"client": dict(db.execute("SELECT * FROM clients WHERE id=?", (existing["id"],)).fetchone()),
                        "updated": True}
            cursor = db.execute("INSERT INTO clients (name, email, rate_cents, currency, notes, created) "
                                "VALUES (:name, :email, :rate_cents, :currency, :notes, :created)", record)
            return {"client": {**record, "id": cursor.lastrowid}, "created": True}
        result = self._run(work)
        client = result["client"]
        client["hourly_rate"] = money(client["rate_cents"], client["currency"])
        return result

    def clients(self) -> dict:
        def work(db):
            rows = db.execute("""SELECT c.*, COALESCE((SELECT SUM(d.total_cents) FROM documents d WHERE d.client_id=c.id
                                 AND d.kind='invoice' AND d.status!='void'), 0) AS invoiced_cents
                                 FROM clients c ORDER BY c.name COLLATE NOCASE""").fetchall()
            return [dict(row) for row in rows]
        clients = self._run(work)
        for client in clients:
            client["hourly_rate"] = money(client["rate_cents"], client["currency"])
            client["invoiced"] = money(client["invoiced_cents"], client["currency"])
        return {"clients": clients, "count": len(clients)}

    def _client(self, db: sqlite3.Connection, name: str) -> sqlite3.Row:
        name = (name or "").strip()
        row = db.execute("SELECT * FROM clients WHERE name=?", (name,)).fetchone()
        if row is None:
            matches = db.execute("SELECT * FROM clients WHERE name LIKE ? ORDER BY name", (f"%{name}%",)).fetchall() if name else []
            if len(matches) == 1:
                return matches[0]
            known = ", ".join(r["name"] for r in db.execute("SELECT name FROM clients ORDER BY name LIMIT 20")) or "none yet"
            raise ValueError(f"Unknown client {name!r}. Add it with client_add first. Known clients: {known}.")
        return row

    # ------------------------------------------------------------------ time
    def log_time(self, client: str, hours: str, description: str, work_date: str | None = None) -> dict:
        minutes = int((_decimal(hours, "Hours", "0.01", "24") * 60).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
        if not (description or "").strip():
            raise ValueError("Describe the work so it can appear on the invoice.")
        day = _day(work_date)

        def work(db):
            row = self._client(db, client)
            cursor = db.execute("INSERT INTO time_entries (client_id, minutes, description, work_date, rate_cents, created) "
                                "VALUES (?, ?, ?, ?, ?, ?)", (row["id"], minutes, description.strip(), day, row["rate_cents"], self._now()))
            value = round(minutes * row["rate_cents"] / 60)
            return {"entry_id": cursor.lastrowid, "client": row["name"], "hours": minutes / 60, "date": day,
                    "value": money(value, row["currency"]),
                    "note": "" if row["rate_cents"] else "This client has no hourly rate yet, so the time is worth 0 until you set one."}
        return self._run(work)

    def start_timer(self, client: str, description: str) -> dict:
        if not (description or "").strip():
            raise ValueError("Say what you are working on.")

        def work(db):
            running = db.execute("SELECT t.*, c.name FROM time_entries t JOIN clients c ON c.id=t.client_id "
                                 "WHERE running_since IS NOT NULL").fetchone()
            if running:
                raise ValueError(f"A timer is already running for {running['name']} since {running['running_since']}. Stop it first.")
            row = self._client(db, client)
            started = self._now()
            db.execute("INSERT INTO time_entries (client_id, minutes, description, work_date, rate_cents, running_since, created) "
                       "VALUES (?, 0, ?, ?, ?, ?, ?)", (row["id"], description.strip(), date.today().isoformat(),
                                                       row["rate_cents"], started, started))
            return {"started": started, "client": row["name"], "description": description.strip()}
        return self._run(work)

    def stop_timer(self) -> dict:
        def work(db):
            running = db.execute("SELECT t.*, c.name, c.currency FROM time_entries t JOIN clients c ON c.id=t.client_id "
                                 "WHERE running_since IS NOT NULL").fetchone()
            if running is None:
                raise ValueError("No timer is running.")
            started = datetime.fromisoformat(running["running_since"])
            minutes = max(1, round((datetime.now(UTC) - started).total_seconds() / 60))
            minutes = min(minutes, 24 * 60)
            db.execute("UPDATE time_entries SET minutes=?, running_since=NULL WHERE id=?", (minutes, running["id"]))
            return {"client": running["name"], "description": running["description"], "hours": round(minutes / 60, 2),
                    "value": money(round(minutes * running["rate_cents"] / 60), running["currency"])}
        return self._run(work)

    def unbilled_time(self) -> dict:
        def work(db):
            return [dict(row) for row in db.execute("""
                SELECT c.name AS client, c.currency, COUNT(*) AS entries, SUM(t.minutes) AS minutes,
                       SUM(t.minutes * t.rate_cents) / 60 AS value_cents
                FROM time_entries t JOIN clients c ON c.id=t.client_id
                WHERE t.invoice_number IS NULL AND t.running_since IS NULL
                GROUP BY c.id ORDER BY c.name""")]
        rows = self._run(work)
        for row in rows:
            row["hours"] = round(row["minutes"] / 60, 2)
            row["value"] = money(int(row["value_cents"] or 0), row["currency"])
        return {"unbilled": rows}

    # ------------------------------------------------------------------ documents
    def _number(self, db: sqlite3.Connection, kind: str, issued: str) -> str:
        prefix = ("INV" if kind == "invoice" else "Q") + f"-{issued[:4]}-"
        last = db.execute("SELECT number FROM documents WHERE number LIKE ? ORDER BY number DESC LIMIT 1",
                          (prefix + "%",)).fetchone()
        return f"{prefix}{(int(last['number'].rsplit('-', 1)[1]) + 1) if last else 1:04d}"

    @staticmethod
    def _items(items: list[dict] | None) -> list[dict]:
        lines = []
        for raw in items or []:
            if not isinstance(raw, dict):
                raise ValueError("Each item needs description, quantity and unit_price.")
            description = str(raw.get("description", "")).strip()
            if not description:
                raise ValueError("Each item needs a description.")
            quantity = _decimal(raw.get("quantity", "1"), "Quantity", "0.01", "100000")
            unit = cents(str(raw.get("unit_price", raw.get("price", ""))))
            amount = int((quantity * unit).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
            lines.append({"description": description, "quantity": format(quantity.normalize(), "f"),
                          "unit_cents": unit, "amount_cents": amount})
        return lines

    def _create_document(self, kind: str, client: str, items: list[dict] | None, days: int, tax_percent: str,
                         currency: str | None, notes: str, include_unbilled_time: bool = False,
                         from_quote: str | None = None) -> dict:
        days = int(days)
        if not 0 <= days <= 365:
            raise ValueError("Days must be between 0 and 365.")
        tax_bp = int(_decimal(tax_percent or "0", "Tax percent", "0", "100") * 100)
        lines = self._items(items)
        issued = date.today().isoformat()

        def work(db):
            row = self._client(db, client)
            doc_currency = _currency(currency, row["currency"])
            billed_ids: list[int] = []
            source = None
            if from_quote:
                quote = db.execute("SELECT * FROM documents WHERE number=? AND kind='quote'", (from_quote.strip().upper(),)).fetchone()
                if quote is None:
                    raise ValueError(f"No quote numbered {from_quote}.")
                if quote["client_id"] != row["id"]:
                    raise ValueError("That quote belongs to a different client.")
                if quote["status"] == "invoiced":
                    raise ValueError(f"Quote {quote['number']} was already turned into an invoice.")
                if quote["currency"] != doc_currency:
                    raise ValueError("The quote's currency differs from this invoice's currency.")
                source = quote["number"]
                lines.extend({"description": item["description"], "quantity": item["quantity"],
                              "unit_cents": item["unit_cents"], "amount_cents": item["amount_cents"]}
                             for item in db.execute("SELECT * FROM document_items WHERE number=? ORDER BY line", (source,)))
                if not tax_bp:
                    nonlocal_tax[0] = quote["tax_bp"]
            if include_unbilled_time:
                if doc_currency != row["currency"]:
                    raise ValueError("Logged time is priced in the client's currency; invoice it in that currency.")
                for entry in db.execute("SELECT * FROM time_entries WHERE client_id=? AND invoice_number IS NULL "
                                        "AND running_since IS NULL ORDER BY work_date, id", (row["id"],)):
                    hours = (Decimal(entry["minutes"]) / 60).quantize(Decimal("0.01"))
                    lines.append({"description": f"{entry['description']} ({entry['work_date']})",
                                  "quantity": format(hours.normalize(), "f"), "unit_cents": entry["rate_cents"],
                                  "amount_cents": round(entry["minutes"] * entry["rate_cents"] / 60)})
                    billed_ids.append(entry["id"])
            if not lines:
                raise ValueError("Nothing to bill: give line items" + (" or log time first." if include_unbilled_time else
                                                                         ", or set include_unbilled_time to bill logged hours."))
            subtotal = sum(line["amount_cents"] for line in lines)
            tax = round(subtotal * nonlocal_tax[0] / 10000)
            number = self._number(db, kind, issued)
            due = (date.today() + timedelta(days=days)).isoformat()
            db.execute("INSERT INTO documents VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, '', ?, ?)",
                       (number, kind, row["id"], doc_currency, issued, due, subtotal, nonlocal_tax[0], tax,
                        subtotal + tax, "unpaid" if kind == "invoice" else "open", (notes or "").strip(), source, self._now()))
            db.executemany("INSERT INTO document_items VALUES (?, ?, ?, ?, ?, ?)",
                           [(number, index, line["description"], line["quantity"], line["unit_cents"], line["amount_cents"])
                            for index, line in enumerate(lines, 1)])
            if billed_ids:
                db.executemany("UPDATE time_entries SET invoice_number=? WHERE id=?", [(number, i) for i in billed_ids])
            if source:
                db.execute("UPDATE documents SET status='invoiced' WHERE number=?", (source,))
            return number

        nonlocal_tax = [tax_bp]
        number = self._run(work)
        try:
            path = self._render(number)
        except Exception:
            # Keep the books and the file consistent: no PDF, no document.
            def undo(db):
                db.execute("UPDATE time_entries SET invoice_number=NULL WHERE invoice_number=?", (number,))
                db.execute("UPDATE documents SET status='open' WHERE number=(SELECT source_quote FROM documents WHERE number=?)", (number,))
                db.execute("DELETE FROM document_items WHERE number=?", (number,))
                db.execute("DELETE FROM documents WHERE number=?", (number,))
            self._run(undo)
            raise
        self._run(lambda db: db.execute("UPDATE documents SET path=? WHERE number=?", (str(path), number)))
        document = self.document(number)
        document["next_step"] = ("Open the PDF and check it, then send it to the client yourself or ask HAYAT to email it "
                                 "(sending asks for your approval).")
        return document

    def create_invoice(self, client: str, items: list[dict] | None = None, due_days: int = 14, tax_percent: str = "0",
                       currency: str | None = None, notes: str = "", include_unbilled_time: bool = False,
                       from_quote: str | None = None) -> dict:
        return self._create_document("invoice", client, items, due_days, tax_percent, currency, notes,
                                     bool(include_unbilled_time), from_quote)

    def create_quote(self, client: str, items: list[dict], valid_days: int = 30, tax_percent: str = "0",
                     currency: str | None = None, notes: str = "") -> dict:
        return self._create_document("quote", client, items, valid_days, tax_percent, currency, notes)

    def document(self, number: str) -> dict:
        number = (number or "").strip().upper()

        def work(db):
            row = db.execute("SELECT d.*, c.name AS client, c.email AS client_email FROM documents d "
                             "JOIN clients c ON c.id=d.client_id WHERE number=?", (number,)).fetchone()
            if row is None:
                raise ValueError(f"No invoice or quote numbered {number}.")
            items = [dict(item) for item in db.execute("SELECT * FROM document_items WHERE number=? ORDER BY line", (number,))]
            paid = db.execute("SELECT COALESCE(SUM(amount_cents), 0) FROM payments WHERE number=?", (number,)).fetchone()[0]
            return dict(row), items, paid
        row, items, paid = self._run(work)
        row["items"] = items
        row["paid_cents"] = paid
        row["balance_cents"] = row["total_cents"] - paid if row["kind"] == "invoice" and row["status"] != "void" else 0
        row["total"] = money(row["total_cents"], row["currency"])
        row["balance"] = money(row["balance_cents"], row["currency"])
        row["overdue"] = bool(row["kind"] == "invoice" and row["status"] == "unpaid" and row["due"] < date.today().isoformat())
        return row

    def invoices(self, status: str | None = None) -> dict:
        wanted = (status or "all").strip().lower()
        if wanted not in {"all", "unpaid", "paid", "overdue", "void", "quotes"}:
            raise ValueError("Status must be all, unpaid, overdue, paid, void or quotes.")
        kind = "quote" if wanted == "quotes" else "invoice"
        numbers = self._run(lambda db: [r[0] for r in db.execute(
            "SELECT number FROM documents WHERE kind=? ORDER BY issued DESC, number DESC", (kind,))])
        documents = [self.document(number) for number in numbers]
        if wanted in {"unpaid", "paid", "void"}:
            documents = [d for d in documents if d["status"] == wanted]
        elif wanted == "overdue":
            documents = [d for d in documents if d["overdue"]]
        summary = [{key: d[key] for key in ("number", "client", "issued", "due", "total", "balance", "status", "overdue", "path",
                                        "currency", "total_cents", "balance_cents")}
                   for d in documents]
        outstanding: dict[str, int] = {}
        for d in documents:
            outstanding[d["currency"]] = outstanding.get(d["currency"], 0) + d["balance_cents"]
        return {"documents": summary, "count": len(summary),
                "outstanding": {cur: money(value, cur) for cur, value in outstanding.items() if value}}

    def record_payment(self, number: str, amount: str | None = None, paid_on: str | None = None,
                       reference: str | None = None) -> dict:
        document = self.document(number)
        if document["kind"] != "invoice":
            raise ValueError("Payments are recorded against invoices, not quotes.")
        if document["status"] == "void":
            raise ValueError("That invoice is void.")
        value = cents(amount) if amount not in (None, "") else document["balance_cents"]
        if value <= 0:
            raise ValueError("Nothing is owed on that invoice." if document["balance_cents"] <= 0 else "Amount must be above 0.")
        if value > document["balance_cents"]:
            raise ValueError(f"That is more than the {document['balance']} still owed.")
        day = _day(paid_on)
        reference = (reference or "").strip() or None

        def work(db):
            if reference and db.execute("SELECT 1 FROM payments WHERE reference=?", (reference,)).fetchone():
                return False
            db.execute("INSERT INTO payments (number, amount_cents, paid_on, reference, recorded) VALUES (?, ?, ?, ?, ?)",
                       (document["number"], value, day, reference, self._now()))
            if value == document["balance_cents"]:
                db.execute("UPDATE documents SET status='paid' WHERE number=?", (document["number"],))
            return True
        recorded = self._run(work)
        after = self.document(document["number"])
        return {"recorded": recorded, "duplicate": not recorded, "number": after["number"], "status": after["status"],
                "paid": money(value, after["currency"]), "balance": after["balance"],
                "source": "owner-reported; HAYAT did not verify the payment"}

    def void_invoice(self, number: str, reason: str) -> dict:
        document = self.document(number)
        if not (reason or "").strip():
            raise ValueError("Give a reason for voiding, for your records.")
        if document["paid_cents"]:
            raise ValueError("Payments were recorded on this invoice; it can't be voided.")

        def work(db):
            db.execute("UPDATE documents SET status='void', notes=TRIM(notes || ' [void: ' || ? || ']') WHERE number=?",
                       (reason.strip(), document["number"]))
            db.execute("UPDATE time_entries SET invoice_number=NULL WHERE invoice_number=?", (document["number"],))
        self._run(work)
        return {"number": document["number"], "status": "void",
                "note": "Any logged time on it is unbilled again. The number is kept so the sequence has no gaps."}

    # ------------------------------------------------------------------ PDF
    def _render(self, number: str) -> Path:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import LETTER
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import inch
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        from xml.sax.saxutils import escape

        from app.tools.filesystem import _pdf_fonts

        document = self.document(number)
        profile = self.profile()
        folder = "invoices" if document["kind"] == "invoice" else "quotes"
        path = self.permissions.validate_project_path(self.output / folder / f"{number}.pdf", write=True)
        path.parent.mkdir(parents=True, exist_ok=True)
        body, bold, unicode_ok = _pdf_fonts()

        def text(value: Any) -> str:
            value = escape(str(value))
            return value if unicode_ok else value.encode("latin-1", "replace").decode("latin-1")

        style = ParagraphStyle("b", fontName=body, fontSize=10, leading=14)
        strong = ParagraphStyle("s", parent=style, fontName=bold)
        title = ParagraphStyle("t", fontName=bold, fontSize=22, leading=28)
        currency = document["currency"]
        heading = "INVOICE" if document["kind"] == "invoice" else "QUOTE"
        seller = "<br/>".join(text(profile[key]) for key in ("business_name", "owner_name", "address", "email") if profile[key])
        flow: list[Any] = [Paragraph(heading, title), Spacer(1, 6)]
        header = Table([[Paragraph(seller or text("(set your business details with business_profile_set)"), style),
                         Paragraph(f"<b>{heading.title()} no.</b> {text(number)}<br/><b>Date</b> {document['issued']}<br/>"
                                   f"<b>{'Due' if document['kind'] == 'invoice' else 'Valid until'}</b> {document['due']}",
                                   ParagraphStyle("r", parent=style, alignment=2))]], colWidths=[3.6 * inch, 2.9 * inch])
        flow += [header, Spacer(1, 14), Paragraph("Bill to", strong),
                 Paragraph(text(document["client"]) + (f"<br/>{text(document['client_email'])}" if document["client_email"] else ""), style),
                 Spacer(1, 14)]
        rows = [[Paragraph("Description", strong), Paragraph("Qty", strong), Paragraph("Unit price", strong), Paragraph("Amount", strong)]]
        rows += [[Paragraph(text(item["description"]), style), item["quantity"], money(item["unit_cents"], currency),
                  money(item["amount_cents"], currency)] for item in document["items"]]
        rows.append(["", "", "Subtotal", money(document["subtotal_cents"], currency)])
        if document["tax_bp"]:
            rows.append(["", "", f"Tax {Decimal(document['tax_bp']) / 100:g}%", money(document["tax_cents"], currency)])
        rows.append(["", "", "Total", money(document["total_cents"], currency)])
        table = Table(rows, colWidths=[3.3 * inch, 0.7 * inch, 1.2 * inch, 1.3 * inch], repeatRows=1)
        table.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), body), ("FONTSIZE", (0, 0), (-1, -1), 10),
            ("LINEBELOW", (0, 0), (-1, 0), 0.8, colors.black), ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
            ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LINEABOVE", (2, -1), (-1, -1), 0.8, colors.black),
            ("FONTNAME", (2, -1), (-1, -1), bold),
        ]))
        flow += [table, Spacer(1, 18)]
        if document["kind"] == "invoice":
            # Pay-now options from the owner's own receiving handles: a UPI QR with the amount filled in, PayPal.me with
            # the amount, payment links. The payer's own app moves the money straight to the owner; HAYAT never does.
            from app.wallet_money import pdf_payment_block
            options = pdf_payment_block(self, document, strong, style, text)
            if profile["payment_instructions"] or options:
                flow += [Paragraph("How to pay", strong)]
                if profile["payment_instructions"]:
                    flow += [Paragraph(text(profile["payment_instructions"]), style)]
                flow += options + [Spacer(1, 10)]
        if document["notes"]:
            flow += [Paragraph("Notes", strong), Paragraph(text(document["notes"]), style)]
        SimpleDocTemplate(str(path), pagesize=LETTER, title=f"{heading.title()} {number}", leftMargin=inch,
                          rightMargin=inch, topMargin=inch, bottomMargin=inch).build(flow)
        return path

    # ------------------------------------------------------------------ expenses
    def record_expense(self, amount: str, category: str, description: str, spent_on: str | None = None,
                       currency: str | None = None, reference: str | None = None) -> dict:
        value = cents(amount)
        if value <= 0:
            raise ValueError("Expense amount must be above 0.")
        if not (category or "").strip() or not (description or "").strip():
            raise ValueError("Category and description are required.")
        record = (value, _currency(currency, self._default_currency()), category.strip().lower(), description.strip(),
                  _day(spent_on), (reference or "").strip() or None, self._now())

        def work(db):
            if record[5] and db.execute("SELECT 1 FROM expenses WHERE reference=?", (record[5],)).fetchone():
                return None
            return db.execute("INSERT INTO expenses (amount_cents, currency, category, description, spent_on, reference, recorded) "
                              "VALUES (?, ?, ?, ?, ?, ?, ?)", record).lastrowid
        expense_id = self._run(work)
        return {"recorded": expense_id is not None, "duplicate": expense_id is None, "id": expense_id,
                "amount": money(value, record[1]), "category": record[2], "date": record[4]}

    def expenses(self, period: str | None = None) -> dict:
        start, end, label = period_bounds(period)
        rows = self._run(lambda db: [dict(r) for r in db.execute(
            "SELECT * FROM expenses WHERE spent_on>=? AND spent_on<? ORDER BY spent_on DESC", (start, end))])
        totals: dict[tuple[str, str], int] = {}
        for row in rows:
            totals[(row["currency"], row["category"])] = totals.get((row["currency"], row["category"]), 0) + row["amount_cents"]
            row["amount"] = money(row["amount_cents"], row["currency"])
        return {"period": label, "expenses": rows,
                "by_category": [{"currency": cur, "category": cat, "total": money(v, cur)} for (cur, cat), v in sorted(totals.items())]}

    # ------------------------------------------------------------------ leads
    def add_lead(self, title: str, source: str, url: str = "", value: str = "0", currency: str | None = None,
                 follow_up: str | None = None, notes: str = "") -> dict:
        if not (title or "").strip() or not (source or "").strip():
            raise ValueError("A lead needs a title and where it came from.")
        if url and not re.match(r"(?i)^https?://", url.strip()):
            raise ValueError("The lead URL must start with http:// or https://.")
        now = self._now()
        record = (title.strip(), source.strip(), (url or "").strip(), cents(value or "0"),
                  _currency(currency, self._default_currency()), "new", _day(follow_up) if follow_up else None,
                  (notes or "").strip(), now, now)
        lead_id = self._run(lambda db: db.execute(
            "INSERT INTO leads (title, source, url, value_cents, currency, status, follow_up, notes, created, updated) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)", record).lastrowid)
        return {"lead_id": lead_id, "title": record[0], "status": "new", "follow_up": record[6]}

    def update_lead(self, lead_id: int, status: str | None = None, follow_up: str | None = None, notes: str | None = None,
                    value: str | None = None) -> dict:
        if status is not None and status.strip().lower() not in LEAD_STATUSES:
            raise ValueError("Status must be one of: " + ", ".join(LEAD_STATUSES) + ".")

        def work(db):
            row = db.execute("SELECT * FROM leads WHERE id=?", (int(lead_id),)).fetchone()
            if row is None:
                raise ValueError(f"No lead with id {lead_id}.")
            new_status = status.strip().lower() if status else row["status"]
            new_follow = (_day(follow_up) if follow_up not in ("", "none", "clear") else None) if follow_up is not None else row["follow_up"]
            if new_status in {"won", "lost"} and follow_up is None:
                new_follow = None
            new_notes = row["notes"] if not notes else (row["notes"] + "\n" if row["notes"] else "") + f"{date.today()}: {notes.strip()}"
            new_value = cents(value) if value not in (None, "") else row["value_cents"]
            db.execute("UPDATE leads SET status=?, follow_up=?, notes=?, value_cents=?, updated=? WHERE id=?",
                       (new_status, new_follow, new_notes, new_value, self._now(), row["id"]))
            return {"lead_id": row["id"], "title": row["title"], "status": new_status, "follow_up": new_follow}
        return self._run(work)

    def contact_record(self, limit: int = 1000) -> str:
        """Every saved lead's name, website and notes, and every client's email, as one text.

        Contact details a reply gives for a lead must come from here (the prospect finder writes
        what each homepage published into the lead's notes), not from the model's memory.
        """
        def work(db):
            rows = db.execute("SELECT title, url, notes FROM leads ORDER BY id DESC LIMIT ?", (limit,)).fetchall()
            clients = db.execute("SELECT name, email, notes FROM clients ORDER BY id DESC LIMIT ?", (limit,)).fetchall()
            return "\n".join(" ".join(str(value or "") for value in row) for row in [*rows, *clients])
        return self._run(work)

    def leads(self, status: str | None = None, due_only: bool = False) -> dict:
        wanted = (status or "open").strip().lower()
        if wanted not in {"open", "all", *LEAD_STATUSES}:
            raise ValueError("Status must be open, all, or one of: " + ", ".join(LEAD_STATUSES) + ".")
        today = date.today().isoformat()

        def work(db):
            query = "SELECT * FROM leads"
            args: list[Any] = []
            clauses = []
            if wanted == "open":
                clauses.append("status NOT IN ('won', 'lost')")
            elif wanted != "all":
                clauses.append("status=?")
                args.append(wanted)
            if due_only:
                clauses.append("follow_up IS NOT NULL AND follow_up<=?")
                args.append(today)
            if clauses:
                query += " WHERE " + " AND ".join(clauses)
            return [dict(r) for r in db.execute(query + " ORDER BY COALESCE(follow_up, '9999'), id", args)]
        rows = self._run(work)
        for row in rows:
            row["value"] = money(row["value_cents"], row["currency"])
            row["follow_up_due"] = bool(row["follow_up"] and row["follow_up"] <= today)
        return {"leads": rows, "count": len(rows), "follow_ups_due": sum(1 for r in rows if r["follow_up_due"])}

    # ------------------------------------------------------------------ goal & report
    def set_goal(self, monthly_amount: str, currency: str | None = None) -> dict:
        value = cents(monthly_amount)
        cur = _currency(currency, self._default_currency())
        self._run(lambda db: db.executemany("INSERT INTO profile VALUES (?, ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
                                            [("goal_cents", str(value)), ("goal_currency", cur)]))
        return {"monthly_goal": money(value, cur)}

    def _goal(self) -> tuple[int, str] | None:
        rows = dict(self._run(lambda db: [(r["key"], r["value"]) for r in db.execute(
            "SELECT key, value FROM profile WHERE key IN ('goal_cents', 'goal_currency')")]))
        if "goal_cents" not in rows or not int(rows["goal_cents"]):
            return None
        return int(rows["goal_cents"]), rows.get("goal_currency", "USD")

    def report(self, period: str | None = None) -> dict:
        start, end, label = period_bounds(period)
        today = date.today().isoformat()

        def work(db):
            income = {r["currency"]: r["total"] for r in db.execute(
                "SELECT d.currency, SUM(p.amount_cents) AS total FROM payments p JOIN documents d ON d.number=p.number "
                "WHERE p.paid_on>=? AND p.paid_on<? GROUP BY d.currency", (start, end))}
            other = {r["currency"]: r["total"] for r in db.execute(
                "SELECT currency, SUM(amount_cents) AS total FROM other_income WHERE received_on>=? AND received_on<? "
                "GROUP BY currency", (start, end))}
            spent = {r["currency"]: r["total"] for r in db.execute(
                "SELECT currency, SUM(amount_cents) AS total FROM expenses WHERE spent_on>=? AND spent_on<? GROUP BY currency",
                (start, end))}
            top_clients = [dict(r) for r in db.execute(
                "SELECT c.name AS client, d.currency, SUM(p.amount_cents) AS paid_cents FROM payments p "
                "JOIN documents d ON d.number=p.number JOIN clients c ON c.id=d.client_id "
                "WHERE p.paid_on>=? AND p.paid_on<? GROUP BY c.id, d.currency ORDER BY paid_cents DESC LIMIT 5", (start, end))]
            return income, other, spent, top_clients
        income, other, spent, top_clients = self._run(work)
        product = self._product_sales(start, end)
        currencies = sorted(set(income) | set(other) | set(spent) | set(product))
        lines = []
        for cur in currencies:
            client_income, product_net, costs = income.get(cur, 0), product.get(cur, 0), spent.get(cur, 0)
            extra = other.get(cur, 0)
            earned = client_income + product_net + extra
            lines.append({"currency": cur, "client_income": money(client_income, cur),
                          "product_sales_net": money(product_net, cur), "other_income": money(extra, cur),
                          "expenses": money(costs, cur),
                          "profit_before_tax": money(earned - costs, cur),
                          "profit_cents": earned - costs,
                          "income_cents": earned})
        unpaid = self.invoices("unpaid")
        overdue = [d for d in unpaid["documents"] if d["overdue"]]
        leads = self.leads("open")
        result: dict[str, Any] = {
            "period": label, "by_currency": lines,
            "top_clients": [{**c, "paid": money(c["paid_cents"], c["currency"])} for c in top_clients],
            "outstanding_invoices": unpaid["outstanding"], "unpaid_count": unpaid["count"],
            "overdue": [{key: d[key] for key in ("number", "client", "due", "balance")} for d in overdue],
            "unbilled_time": self.unbilled_time()["unbilled"],
            "open_leads": leads["count"], "follow_ups_due": [
                {"id": lead["id"], "title": lead["title"], "follow_up": lead["follow_up"]}
                for lead in leads["leads"] if lead["follow_up_due"]],
            "basis": BASIS,
        }
        goal = self._goal()
        if goal and re.fullmatch(r"\d{4}-\d{2}-01", start) and end != "9999-12-31" and end[:7] != start[:7] and \
                (date.fromisoformat(end) - date.fromisoformat(start)).days <= 31:
            target, cur = goal
            earned = next((line["income_cents"] for line in lines if line["currency"] == cur), 0)
            result["goal"] = {"target": money(target, cur), "earned": money(earned, cur),
                              "percent": round(100 * earned / target, 1), "remaining": money(max(0, target - earned), cur)}
            if start <= today < end:
                first = date.fromisoformat(start)
                days_total = (date.fromisoformat(end) - first).days
                elapsed = (date.today() - first).days + 1
                result["goal"]["on_pace"] = earned * days_total >= target * elapsed
        balances = self._balances()
        result["wallet"] = {cur: money(row["available"], cur) for cur, row in balances.items() if row["available"]}
        result["payout_ready"] = self.payout()["configured"]
        result["next_actions"] = self._next_actions(result)
        return result

    @staticmethod
    def _next_actions(report: dict) -> list[str]:
        actions = []
        if report["overdue"]:
            actions.append(f"Chase {len(report['overdue'])} overdue invoice(s): " + ", ".join(d["number"] for d in report["overdue"][:5]) + ".")
        for item in report["unbilled_time"]:
            hours = f"{item['hours']:g}"
            actions.append(f"Invoice {hours} unbilled hour{'' if hours == '1' else 's'} for {item['client']} ({item['value']}).")
        if report["follow_ups_due"]:
            actions.append(f"Follow up on {len(report['follow_ups_due'])} lead(s) due today or earlier.")
        goal = report.get("goal")
        if goal and goal.get("on_pace") is False:
            actions.append(f"Behind the monthly goal: {goal['remaining']} still to earn.")
        if not report["by_currency"] and not report["open_leads"]:
            actions.append("No income or leads recorded yet. Ask HAYAT to research work you could sell, then add leads.")
        if report.get("wallet") and not report.get("payout_ready"):
            actions.append(f"{', '.join(report['wallet'].values())} is in your wallet. "
                           "Set up a withdrawal method whenever you are ready.")
        return actions

    # ------------------------------------------------------------------ wallet
    #: Where earnings can be withdrawn to. HAYAT only records the choice; the
    #: transfer itself happens in that account.
    PAYOUT_METHODS = ("bank", "paypal", "stripe", "wise", "payoneer", "upi", "other")

    def record_income(self, amount: str, source: str, description: str, received_on: str | None = None,
                      currency: str | None = None, reference: str | None = None) -> dict:
        """Money earned outside invoices and the product store: a platform payout, a tip, a sale elsewhere."""
        value = cents(amount)
        if value <= 0:
            raise ValueError("Income must be above 0.")
        if not (source or "").strip() or not (description or "").strip():
            raise ValueError("Say where the money came from and what it was for.")
        record = (value, _currency(currency, self._default_currency()), source.strip(), description.strip(),
                  _day(received_on), (reference or "").strip() or None, self._now())

        def work(db):
            if record[5] and db.execute("SELECT 1 FROM other_income WHERE reference=?", (record[5],)).fetchone():
                return None
            return db.execute("INSERT INTO other_income (amount_cents, currency, source, description, received_on, "
                              "reference, recorded) VALUES (?, ?, ?, ?, ?, ?, ?)", record).lastrowid
        income_id = self._run(work)
        return {"recorded": income_id is not None, "duplicate": income_id is None, "id": income_id,
                "amount": money(value, record[1]), "source": record[2], "date": record[4],
                "wallet": self.wallet(activity=0)["balances"]}

    def payout(self) -> dict:
        profile = dict(self._run(lambda db: [(r["key"], r["value"]) for r in db.execute(
            "SELECT key, value FROM profile WHERE key LIKE 'payout_%'")]))
        if not profile.get("payout_method"):
            return {"configured": False,
                    "note": "No withdrawal method yet. Set one up whenever you are ready; earnings keep adding up meanwhile."}
        last4 = profile.get("payout_last4", "")
        shown = profile["payout_name"] + (f" ending {last4}" if last4 else "")
        return {"configured": True, "method": profile["payout_method"], "name": profile["payout_name"],
                "last4": last4, "currency": profile.get("payout_currency", ""), "set_on": profile.get("payout_set_on", ""),
                "destination": f"{self.METHOD_NAMES[profile['payout_method']]}: {shown}"}

    METHOD_NAMES = {"bank": "Bank", "paypal": "PayPal", "stripe": "Stripe", "wise": "Wise",
                    "payoneer": "Payoneer", "upi": "UPI", "other": "Other"}

    def set_payout(self, method: str, name: str, last4: str = "", currency: str | None = None) -> dict:
        """Remember where withdrawals go, without anything that could move money.

        Only a method, a name the owner recognises and the last four digits
        are kept. A full account or card number, a password or a PIN is
        refused: those belong in the payment provider, never in HAYAT.
        """
        method = (method or "").strip().lower()
        if method not in self.PAYOUT_METHODS:
            raise ValueError("Method must be one of: " + ", ".join(self.PAYOUT_METHODS) + ".")
        name = (name or "").strip()
        if not name or len(name) > 80:
            raise ValueError("Give the account a short name you recognise, e.g. 'My savings account' or your PayPal email.")
        if re.search(r"\d{6,}", name.replace(" ", "").replace("-", "")):
            raise ValueError("Don't store account or card numbers. Use a name, plus only the last 4 digits.")
        if re.search(r"(?i)\b(?:password|passcode|pin|cvv|cvc|otp)\b", name):
            raise ValueError("Never give HAYAT passwords, PINs or security codes.")
        last4 = (last4 or "").strip()
        if last4 and not re.fullmatch(r"\d{4}", last4):
            raise ValueError("Only the last 4 digits, or leave it empty.")
        cur = _currency(currency, self._default_currency())
        values = {"payout_method": method, "payout_name": name, "payout_last4": last4, "payout_currency": cur,
                  "payout_set_on": date.today().isoformat()}
        self._run(lambda db: db.executemany(
            "INSERT INTO profile VALUES (?, ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", list(values.items())))
        return {**self.payout(), "note": "Saved. HAYAT records withdrawals; you make each transfer in that account."}

    def clear_payout(self) -> dict:
        self._run(lambda db: db.execute("DELETE FROM profile WHERE key LIKE 'payout_%'"))
        return self.payout()

    def _balances(self) -> dict[str, dict[str, int]]:
        """Per currency, all time: earned, spent, withdrawn, pending and available."""
        def work(db):
            rows: dict[str, dict[str, int]] = {}

            def add(currency: str, key: str, value: int) -> None:
                rows.setdefault(currency, {"earned": 0, "funded": 0, "expenses": 0, "withdrawn": 0, "pending": 0})[key] += int(value or 0)
            for r in db.execute("SELECT d.currency, SUM(p.amount_cents) FROM payments p "
                                "JOIN documents d ON d.number=p.number GROUP BY d.currency"):
                add(r[0], "earned", r[1])
            for r in db.execute("SELECT currency, SUM(amount_cents) FROM other_income GROUP BY currency"):
                add(r[0], "earned", r[1])
            for r in db.execute("SELECT currency, SUM(amount_cents) FROM expenses GROUP BY currency"):
                add(r[0], "expenses", r[1])
            for r in db.execute("SELECT currency, SUM(amount_cents) FROM funding GROUP BY currency"):
                add(r[0], "funded", r[1])
            for r in db.execute("SELECT currency, status, SUM(amount_cents) FROM withdrawals "
                                "WHERE status IN ('pending', 'done') GROUP BY currency, status"):
                add(r[0], "withdrawn" if r[1] == "done" else "pending", r[2])
            return rows
        rows = self._run(work)
        for cur, value in self._product_sales("0000-01-01", "9999-12-31").items():
            rows.setdefault(cur, {"earned": 0, "funded": 0, "expenses": 0, "withdrawn": 0, "pending": 0})["earned"] += value
        for row in rows.values():
            row["available"] = row["earned"] + row["funded"] - row["expenses"] - row["withdrawn"] - row["pending"]
        return dict(sorted(rows.items()))

    def wallet(self, activity: int = 25) -> dict:
        """The earnings wallet: what is available to withdraw, and how it got there."""
        balances = self._balances()
        result: dict[str, Any] = {
            "balances": [{"currency": cur, "earned": money(row["earned"], cur), "added": money(row["funded"], cur),
                          "added_cents": row["funded"], "expenses": money(row["expenses"], cur),
                          "withdrawn": money(row["withdrawn"], cur), "pending": money(row["pending"], cur),
                          "available": money(row["available"], cur), "available_cents": row["available"]}
                         for cur, row in balances.items()],
            "payout": self.payout(),
            "basis": ("HAYAT keeps the record, not the money: your earnings sit in the bank, PayPal or store "
                      "account they were paid into. Available = earned + money you added - expenses - withdrawals. "
                      "Money you added is your own, recorded as you said; it is never counted as income."),
        }
        if activity:
            result["activity"] = self._activity(activity)
            result["withdrawals"] = self.withdrawals()
        return result

    def _activity(self, limit: int) -> list[dict]:
        def work(db):
            rows = [("in", r[0], "Invoice payment", f"{r[1]} from {r[2]}", r[3], r[4]) for r in db.execute(
                "SELECT p.paid_on, p.number, c.name, d.currency, p.amount_cents FROM payments p "
                "JOIN documents d ON d.number=p.number JOIN clients c ON c.id=d.client_id")]
            rows += [("in", r[0], "Income", f"{r[1]}: {r[2]}", r[3], r[4]) for r in db.execute(
                "SELECT received_on, source, description, currency, amount_cents FROM other_income")]
            rows += [("out", r[0], "Expense", f"{r[1]}: {r[2]}", r[3], -r[4]) for r in db.execute(
                "SELECT spent_on, category, description, currency, amount_cents FROM expenses")]
            rows += [("in", r[0], "Money you added", (f"from {r[1]}" if r[1] else "your own money") + (f": {r[2]}" if r[2] else ""),
                      r[3], r[4]) for r in db.execute(
                "SELECT f.added_on, m.nickname, f.note, f.currency, f.amount_cents FROM funding f "
                "LEFT JOIN money_methods m ON m.id=f.method_id")]
            rows += [("out", r[0], "Withdrawal" + ("" if r[5] == "done" else f" ({r[5]})"), f"#{r[1]} to {r[2]}",
                      r[3], -r[4] if r[5] != "cancelled" else 0) for r in db.execute(
                "SELECT COALESCE(completed_on, requested_on), id, destination, currency, amount_cents, status "
                "FROM withdrawals")]
            return rows
        rows = self._run(work)
        if self.products is not None:
            try:
                rows += [("in", sale["date"], "Product sale", sale["title"], sale["currency"], sale["net_cents"])
                         for sale in self.products.sales_list()]
            except (sqlite3.Error, AttributeError):
                pass
        rows.sort(key=lambda row: row[1], reverse=True)
        return [{"date": day[:10], "kind": kind, "description": text, "currency": cur, "amount_cents": value,
                 "amount": ("+" if value > 0 else "") + money(value, cur), "direction": direction}
                for direction, day, kind, text, cur, value in rows[:limit]]

    def withdraw(self, amount: str, currency: str | None = None, note: str = "") -> dict:
        """Record a withdrawal to the set-up method. HAYAT moves no money."""
        payout = self.payout()
        if not payout["configured"]:
            raise ValueError("Set up a withdrawal method first (wallet_payout_setup). Your earnings stay recorded meanwhile.")
        cur = _currency(currency, payout["currency"] or self._default_currency())
        available = self._balances().get(cur, {}).get("available", 0)
        value = available if str(amount).strip().lower() in {"all", "everything", "max"} else cents(amount)
        if value <= 0:
            raise ValueError(f"Nothing available to withdraw in {cur}." if available <= 0 else "Amount must be above 0.")
        if value > available:
            raise ValueError(f"Only {money(available, cur)} is available to withdraw.")
        withdrawal_id = self._run(lambda db: db.execute(
            "INSERT INTO withdrawals (amount_cents, currency, destination, status, note, requested_on, recorded) "
            "VALUES (?, ?, ?, 'pending', ?, ?, ?)",
            (value, cur, payout["destination"], (note or "").strip(), date.today().isoformat(), self._now())).lastrowid)
        return {"withdrawal_id": withdrawal_id, "amount": money(value, cur), "to": payout["destination"],
                "status": "pending", "available_after": money(available - value, cur),
                "next_step": (f"HAYAT does not move money. Transfer {money(value, cur)} in your {self.METHOD_NAMES[payout['method']]} "
                              f"account, then say 'withdrawal {withdrawal_id} is done'.")}

    def complete_withdrawal(self, withdrawal_id: int, reference: str = "", completed_on: str | None = None) -> dict:
        day = _day(completed_on)

        def work(db):
            row = db.execute("SELECT * FROM withdrawals WHERE id=?", (int(withdrawal_id),)).fetchone()
            if row is None:
                raise ValueError(f"No withdrawal #{withdrawal_id}.")
            if row["status"] != "pending":
                raise ValueError(f"Withdrawal #{withdrawal_id} is already {row['status']}.")
            db.execute("UPDATE withdrawals SET status='done', completed_on=?, reference=? WHERE id=?",
                       (day, (reference or "").strip() or None, row["id"]))
            return dict(row)
        row = self._run(work)
        return {"withdrawal_id": row["id"], "status": "done", "amount": money(row["amount_cents"], row["currency"]),
                "source": "owner-reported; HAYAT did not see the transfer"}

    def cancel_withdrawal(self, withdrawal_id: int) -> dict:
        def work(db):
            row = db.execute("SELECT * FROM withdrawals WHERE id=?", (int(withdrawal_id),)).fetchone()
            if row is None:
                raise ValueError(f"No withdrawal #{withdrawal_id}.")
            if row["status"] != "pending":
                raise ValueError(f"Withdrawal #{withdrawal_id} is {row['status']}; only a pending one can be cancelled.")
            db.execute("UPDATE withdrawals SET status='cancelled' WHERE id=?", (row["id"],))
            return dict(row)
        row = self._run(work)
        return {"withdrawal_id": row["id"], "status": "cancelled",
                "note": f"{money(row['amount_cents'], row['currency'])} is available again."}

    def withdrawals(self) -> list[dict]:
        rows = self._run(lambda db: [dict(r) for r in db.execute("SELECT * FROM withdrawals ORDER BY id DESC")])
        for row in rows:
            row["amount"] = money(row["amount_cents"], row["currency"])
        return rows

    # ------------------------------------------------------------------ income streams
    def _stream_rows(self, start: str, end: str) -> list[dict]:
        def work(db):
            rows = [{"kind": "Client work", "source": r["name"], "currency": r["currency"], "cents": r["total"],
                     "entries": r["entries"]} for r in db.execute(
                "SELECT c.name, d.currency, SUM(p.amount_cents) AS total, COUNT(*) AS entries FROM payments p "
                "JOIN documents d ON d.number=p.number JOIN clients c ON c.id=d.client_id "
                "WHERE p.paid_on>=? AND p.paid_on<? GROUP BY c.id, d.currency", (start, end))]
            rows += [{"kind": "Other income", "source": r["source"], "currency": r["currency"], "cents": r["total"],
                      "entries": r["entries"]} for r in db.execute(
                "SELECT source, currency, SUM(amount_cents) AS total, COUNT(*) AS entries FROM other_income "
                "WHERE received_on>=? AND received_on<? GROUP BY source COLLATE NOCASE, currency", (start, end))]
            return rows
        rows = self._run(work)
        if self.products is not None:
            try:
                rows += [{"kind": "Digital products", "source": r["title"], "currency": r["currency"],
                          "cents": int(r["cents"] or 0), "entries": r["entries"]}
                         for r in self.products.net_by_product(start, end)]
            except (sqlite3.Error, AttributeError):
                pass
        return rows

    def income_streams(self, period: str | None = None) -> dict:
        """Where the money actually comes from: every recorded source, its share, and its change.

        Built only from what was recorded (invoice payments, product sales,
        other income). It estimates nothing and predicts nothing.
        """
        start, end, label = period_bounds(period)
        rows = self._stream_rows(start, end)
        previous: dict[tuple[str, str, str], int] = {}
        if re.fullmatch(r"\d{4}-\d{2}-01", start) and end != "9999-12-31":
            first = date.fromisoformat(start)
            before = (first - timedelta(days=1)).replace(day=1)
            previous = {(r["kind"], r["source"].lower(), r["currency"]): r["cents"]
                        for r in self._stream_rows(before.isoformat(), start)}
        totals: dict[str, int] = {}
        for row in rows:
            totals[row["currency"]] = totals.get(row["currency"], 0) + row["cents"]
        streams = []
        for row in sorted(rows, key=lambda r: (r["currency"], -r["cents"])):
            total = totals[row["currency"]] or 1
            last = previous.get((row["kind"], row["source"].lower(), row["currency"])) if previous else None
            streams.append({**row, "amount": money(row["cents"], row["currency"]),
                            "share": round(100 * row["cents"] / total, 1),
                            "change": None if last is None else money(row["cents"] - last, row["currency"]),
                            "new": bool(previous) and last is None})
        kinds = sorted({row["kind"] for row in rows})
        notes = []
        for currency, total in totals.items():
            top = max((s for s in streams if s["currency"] == currency), key=lambda s: s["cents"])
            if total and top["share"] >= 70:
                notes.append(f"{top['share']:g}% of your {currency} income came from {top['source']} "
                             f"({top['kind'].lower()}). A second stream would make you less dependent on it.")
        if not rows:
            notes.append("No income recorded for this period. Record payments, product sales or other income and "
                         "each source appears here.")
        return {"period": label, "streams": streams, "kinds": kinds, "stream_count": len(streams),
                "totals": {currency: money(value, currency) for currency, value in totals.items()},
                "notes": notes, "basis": "Only what you recorded. Nothing here is estimated or predicted."}

    def _product_sales(self, start: str, end: str) -> dict[str, int]:
        if self.products is None:
            return {}
        try:
            return self.products.net_by_currency(start, end)
        except (sqlite3.Error, AttributeError):
            return {}

    def export(self, year: str | None = None) -> dict:
        from openpyxl import Workbook

        year = str(year or date.today().year).strip()
        start, end, label = period_bounds(year)
        if not re.fullmatch(r"\d{4}", year):
            raise ValueError("Export one calendar year at a time, e.g. 2026.")

        def work(db):
            payments = db.execute("SELECT p.paid_on, p.number, c.name, d.currency, p.amount_cents, COALESCE(p.reference, '') "
                                  "FROM payments p JOIN documents d ON d.number=p.number JOIN clients c ON c.id=d.client_id "
                                  "WHERE p.paid_on>=? AND p.paid_on<? ORDER BY p.paid_on", (start, end)).fetchall()
            expenses = db.execute("SELECT spent_on, category, description, currency, amount_cents, COALESCE(reference, '') "
                                  "FROM expenses WHERE spent_on>=? AND spent_on<? ORDER BY spent_on", (start, end)).fetchall()
            invoices = db.execute("SELECT d.number, d.issued, d.due, c.name, d.currency, d.subtotal_cents, d.tax_cents, "
                                  "d.total_cents, d.status FROM documents d JOIN clients c ON c.id=d.client_id "
                                  "WHERE d.kind='invoice' AND d.issued>=? AND d.issued<? ORDER BY d.number", (start, end)).fetchall()
            other = db.execute("SELECT received_on, source, description, currency, amount_cents, COALESCE(reference, '') "
                               "FROM other_income WHERE received_on>=? AND received_on<? ORDER BY received_on",
                               (start, end)).fetchall()
            withdrawals = db.execute("SELECT requested_on, COALESCE(completed_on, ''), destination, currency, amount_cents, "
                                     "status, COALESCE(reference, '') FROM withdrawals "
                                     "WHERE requested_on>=? AND requested_on<? ORDER BY id", (start, end)).fetchall()
            return payments, expenses, invoices, other, withdrawals
        payments, expenses, invoices, other, withdrawals = self._run(work)
        book = Workbook()
        sheets = (
            ("Income received", ("Date", "Invoice", "Client", "Currency", "Amount", "Reference"), payments, (4,)),
            ("Expenses", ("Date", "Category", "Description", "Currency", "Amount", "Reference"), expenses, (4,)),
            ("Invoices issued", ("Number", "Issued", "Due", "Client", "Currency", "Subtotal", "Tax", "Total", "Status"),
             invoices, (5, 6, 7)),
            ("Other income", ("Date", "Source", "Description", "Currency", "Amount", "Reference"), other, (4,)),
            ("Withdrawals", ("Requested", "Completed", "To", "Currency", "Amount", "Status", "Reference"), withdrawals, (4,)),
        )
        for index, (name, header, rows, money_cols) in enumerate(sheets):
            sheet = book.active if index == 0 else book.create_sheet()
            sheet.title = name
            sheet.append(list(header))
            for row in rows:
                values = list(row)
                for col in money_cols:
                    values[col] = values[col] / 100
                sheet.append(values)
        product_sheet = book.create_sheet("Product sales")
        product_sheet.append(["Currency", "Net receipts"])
        for cur, value in sorted(self._product_sales(start, end).items()):
            product_sheet.append([cur, value / 100])
        notes = book.create_sheet("Read me")
        for line in (f"HAYAT business records for {label}.", BASIS,
                     "Check every figure against your bank and store statements before filing taxes."):
            notes.append([line])
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        path = self.permissions.validate_project_path(self.output / "exports" / f"business-{year}-{stamp}.xlsx", write=True)
        path.parent.mkdir(parents=True, exist_ok=True)
        book.save(path)
        return {"path": str(path), "year": year, "income_rows": len(payments), "expense_rows": len(expenses),
                "invoice_rows": len(invoices), "note": "Check it against bank and store statements before filing taxes."}


def register_business_tools(registry: ToolRegistry, ledger: BusinessLedger) -> None:
    def schema(properties: dict, required: tuple = ()) -> dict:
        return {"type": "object", "properties": properties, "required": list(required), "additionalProperties": False}

    text = {"type": "string", "minLength": 1, "maxLength": 4000}
    optional = {"type": "string", "maxLength": 4000}
    item = {"type": "object", "properties": {"description": text, "quantity": {"type": ["string", "number"]},
                                               "unit_price": {"type": ["string", "number"]}},
            "required": ["description", "unit_price"], "additionalProperties": False}
    items = {"type": "array", "items": item, "maxItems": 200}
    write, read = "filesystem_write_project", "filesystem_read"

    def reg(name: str, description: str, properties: dict, required: tuple, category: str, handler) -> None:
        registry.register(name, description, schema(properties, required), "low", category, handler)

    reg("business_profile_set", "Save the owner's business name, contact details, address, default currency and how clients should pay "
        "(shown on invoices). Never store passwords or card numbers here.",
        {key: optional for key in BusinessLedger.PROFILE_KEYS}, (), write, ledger.set_profile)
    reg("business_profile", "Show the saved business details used on invoices.", {}, (), read, ledger.profile)
    from app.business_model import DEFAULTS, NUMERIC_FIELDS
    assumptions = schema({key: {"type": ["string", "number"]} if key in NUMERIC_FIELDS else optional for key in DEFAULTS})
    reg("business_model", "Read or evaluate monthly pricing, unit costs, owner time, demand, capacity, break-even, downside/base/upside scenarios, one-factor sensitivity and target price/conversion/acquisition-cost thresholds. Estimates are separate from accounting.",
        {"assumptions": assumptions}, (), read, ledger.business_model)
    reg("business_model_save", "Save a business planning scenario locally. Does not record income or alter invoices. Fields: customer, problem, offer, channel, advantage, experiment, currency, price, direct_cost, acquisition_cost, hours_per_sale, hourly_value, capacity_hours, monthly_leads, conversion_percent, fixed_cost, target_profit.",
        {"assumptions": assumptions}, ("assumptions",), write, ledger.save_business_model)
    from app.business_pilot import PILOT_FIELDS
    observations = schema({**{key: {'type': ['string', 'number']} for key in PILOT_FIELDS}, 'notes': optional})
    reg('business_pilot', 'Read the last pilot check, or compare owner-reported cohort leads, paying customers, net received revenue, total direct costs, acquisition spend and delivery hours against a planning scenario. Does not verify demand or change accounting.',
        {'observations': observations, 'assumptions': assumptions}, (), read, ledger.business_pilot)
    reg('business_pilot_save', 'Save a local pilot comparison with the exact planning assumptions and timestamp. Costs and hours are totals for the same cohort. Does not record revenue or payments in accounting.',
        {'observations': observations, 'assumptions': assumptions}, ('observations',), write, ledger.save_business_pilot)
    reg("client_add", "Add a client (or update an existing one) with email, hourly rate and currency.",
        {"name": text, "email": optional, "hourly_rate": optional, "currency": optional, "notes": optional}, ("name",),
        write, ledger.add_client)
    reg("client_list", "List clients with their hourly rates and total invoiced.", {}, (), read, ledger.clients)
    reg("time_log", "Log billable hours worked for a client (decimal hours, e.g. 1.5). Logged time can be billed later.",
        {"client": text, "hours": {"type": ["string", "number"]}, "description": text, "work_date": optional},
        ("client", "hours", "description"), write, ledger.log_time)
    reg("time_timer_start", "Start a work timer for a client. Only one timer runs at a time.",
        {"client": text, "description": text}, ("client", "description"), write, ledger.start_timer)
    reg("time_timer_stop", "Stop the running work timer and save the time worked.", {}, (), write, ledger.stop_timer)
    reg("time_unbilled", "Show logged hours not yet invoiced, per client, with their value.", {}, (), read, ledger.unbilled_time)
    reg("quote_create", "Create a numbered quote/estimate PDF for a client from line items. Does not send it.",
        {"client": text, "items": items, "valid_days": {"type": "integer", "minimum": 0, "maximum": 365},
         "tax_percent": optional, "currency": optional, "notes": optional}, ("client", "items"), write, ledger.create_quote)
    reg("invoice_create", "Create a numbered invoice PDF for a client from line items, logged unbilled time "
        "(include_unbilled_time), or an accepted quote (from_quote). Does not send it.",
        {"client": text, "items": items, "due_days": {"type": "integer", "minimum": 0, "maximum": 365},
         "tax_percent": optional, "currency": optional, "notes": optional,
         "include_unbilled_time": {"type": "boolean"}, "from_quote": optional}, ("client",), write, ledger.create_invoice)
    reg("invoice_list", "List invoices (status: all, unpaid, overdue, paid, void) or quotes (status: quotes) with balances.",
        {"status": optional}, (), read, ledger.invoices)
    reg("invoice_get", "Show one invoice or quote with its line items, payments, balance and PDF path.",
        {"number": text}, ("number",), read, ledger.document)
    reg("invoice_record_payment", "Record money the owner says was received for an invoice (full balance if no amount). "
        "A repeated reference is not counted twice.",
        {"number": text, "amount": optional, "paid_on": optional, "reference": optional}, ("number",), write, ledger.record_payment)
    reg("invoice_void", "Void an unpaid invoice that was issued by mistake, with a reason. Its logged time becomes unbilled again.",
        {"number": text, "reason": text}, ("number", "reason"), write, ledger.void_invoice)
    reg("expense_record", "Record a business expense (amount, category such as software, equipment, ads, fees, travel).",
        {"amount": {"type": ["string", "number"]}, "category": text, "description": text, "spent_on": optional,
         "currency": optional, "reference": optional}, ("amount", "category", "description"), write, ledger.record_expense)
    reg("expense_list", "List expenses for a period with totals by category.", {"period": optional}, (), read, ledger.expenses)
    reg("lead_add", "Track a work opportunity (a job post, client inquiry, gig or contract) with its source, link, value and "
        "follow-up date.",
        {"title": text, "source": text, "url": optional, "value": optional, "currency": optional,
         "follow_up": optional, "notes": optional}, ("title", "source"), write, ledger.add_lead)
    reg("lead_update", "Update a lead's status (new, contacted, proposal, won, lost), follow-up date, value or notes.",
        {"lead_id": {"type": "integer", "minimum": 1}, "status": optional, "follow_up": optional, "notes": optional,
         "value": optional}, ("lead_id",), write, ledger.update_lead)
    reg("lead_list", "List work leads (status: open, all, new, contacted, proposal, won, lost); due_only shows follow-ups due.",
        {"status": optional, "due_only": {"type": "boolean"}}, (), read, ledger.leads)
    reg("income_goal_set", "Set a monthly income goal that the business report tracks.",
        {"monthly_amount": {"type": ["string", "number"]}, "currency": optional}, ("monthly_amount",), write, ledger.set_goal)
    reg("business_report", "Income, expenses and profit before tax for a period (this month, last month, this year, YYYY-MM, "
        "YYYY, all), including product sales, unpaid and overdue invoices, unbilled time, lead follow-ups, goal progress "
        "and next actions.", {"period": optional}, (), read, ledger.report)
    reg("income_streams", "Where the money actually comes from: every recorded income source (each client, product "
        "and other source such as a marketplace), its amount, share and change since last period, from the books only. "
        "Period: this month, last month, this year, YYYY-MM, YYYY, all.", {"period": optional}, (), read,
        ledger.income_streams)
    reg("business_export", "Export a year's income, expenses and invoices to an Excel workbook for tax preparation.",
        {"year": optional}, (), write, ledger.export)
    # The wallet. These record and report; none of them moves money, and the
    # payout setup refuses full account numbers, passwords and PINs.
    reg("wallet_balance", "Show the earnings wallet: money available to withdraw per currency (earned from invoices, "
        "product sales and other income, minus expenses and withdrawals), the withdrawal method if set up, "
        "recent activity and withdrawals.", {}, (), read, ledger.wallet)
    reg("wallet_record_income", "Add money earned outside invoices and the product store to the wallet, e.g. a "
        "marketplace payout, a tip or a sale on another site.",
        {"amount": {"type": ["string", "number"]}, "source": text, "description": text, "received_on": optional,
         "currency": optional, "reference": optional}, ("amount", "source", "description"), write, ledger.record_income)
    reg("wallet_payout_setup", "Set up where withdrawals go: method (bank, paypal, stripe, wise, payoneer, upi, other), "
        "a recognisable account name and optionally the last 4 digits. Never pass full account or card numbers, "
        "passwords or PINs.",
        {"method": text, "name": text, "last4": optional, "currency": optional}, ("method", "name"), write, ledger.set_payout)
    reg("wallet_payout_clear", "Remove the saved withdrawal method.", {}, (), write, ledger.clear_payout)
    reg("wallet_withdraw", "Record a withdrawal from the wallet to the set-up method (amount, or 'all'). It stays pending "
        "until the owner says the transfer is done. HAYAT does not move money.",
        {"amount": {"type": ["string", "number"]}, "currency": optional, "note": optional}, ("amount",), write,
        ledger.withdraw)
    reg("wallet_withdrawal_done", "Mark a pending withdrawal as transferred, with an optional bank or PayPal reference.",
        {"withdrawal_id": {"type": "integer", "minimum": 1}, "reference": optional, "completed_on": optional},
        ("withdrawal_id",), write, ledger.complete_withdrawal)
    reg("wallet_withdrawal_cancel", "Cancel a pending withdrawal; its amount becomes available again.",
        {"withdrawal_id": {"type": "integer", "minimum": 1}}, ("withdrawal_id",), write, ledger.cancel_withdrawal)
